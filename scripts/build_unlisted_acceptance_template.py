"""生成不在已核验模板白名单中的问题库闭环验收模板。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def build(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.6)
    section.left_margin = section.right_margin = Cm(2.8)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(9)
    for name in ("Title", "Heading 1"):
        style = document.styles[name]
        style.font.name = "黑体"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
    title_border = document.styles["Title"]._element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        document.styles["Title"]._element.pPr.remove(title_border)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("2026年第一期地类变化问数验收快报")
    title.runs[0].font.name = "黑体"
    title.runs[0].font.size = Pt(18)
    title.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    document.add_paragraph("查询时间：[[SYSTEM_TIME]]")
    document.add_paragraph(
        "本快报汇总2026年第一期地类变化图斑的总体规模、耕地及建设用地变化，"
        "并列出需结合审批信息进一步核实的事项。面积统一以亩表示。"
    )

    def heading(text: str) -> None:
        paragraph = document.add_heading(text, level=1)
        paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    def table(headers: list[str], rows: int = 5) -> None:
        grid = document.add_table(rows=rows + 1, cols=len(headers))
        grid.alignment = WD_TABLE_ALIGNMENT.CENTER
        grid.autofit = False
        for row_index, row in enumerate(grid.rows):
            for column, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = headers[column] if row_index == 0 else ""
                cell.width = Cm(4.5 if column == 0 else 2.8)
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:val"), "single")
                    edge.set(qn("w:sz"), "5")
                    edge.set(qn("w:color"), "D9D9D9")
                    borders.append(edge)
                properties.append(borders)
                if row_index == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "E7E9EC")
                    properties.append(fill)

    heading("一 全省变化概况")
    document.add_paragraph(
        "统计口径为全库共有多少个变化图斑、总面积是多少。本期全库共有"
        "{{total_summary.n}}个变化图斑，总面积{{total_summary.area|mu}}亩。"
    )
    document.add_paragraph(
        "按变化后图斑类型汇总面积，主要类别及其图斑数量见表1；"
        "面积构成见图1。"
    )
    document.add_paragraph("表1 变化后图斑类型面积统计")
    document.add_paragraph("单位：亩")
    table(["变化后图斑类型", "图斑数量", "面积"], rows=6)
    document.add_paragraph("图1 变化后主要图斑类型面积构成")

    heading("二 耕地变化情况")
    document.add_paragraph(
        "筛选出所有耕地图斑，统计其图斑数量和总面积；此处指变化前为耕地的图斑，"
        "共{{orig_farmland.n}}个，面积{{orig_farmland.area_m2|mu}}亩。"
    )
    document.add_paragraph(
        "统计变化后为耕地的图斑数量和总面积：共{{current_farmland.n}}个，"
        "面积{{current_farmland.area_m2|mu}}亩。"
    )
    document.add_paragraph(
        "原耕地流向建设用地的图斑共{{farmland_to_construction.n}}个，"
        "面积{{farmland_to_construction.area|mu}}亩；"
        "这一统计不等同于违法占用耕地。"
    )
    document.add_paragraph("分县变化规模及排序见表2，面积对比见图2。")
    document.add_paragraph("表2 各县变化图斑面积排序")
    document.add_paragraph("单位：亩")
    table(["县（市、区）", "图斑数量", "面积"], rows=6)
    document.add_paragraph("图2 各县变化图斑面积对比")

    heading("三 建设用地及图斑特征")
    document.add_paragraph(
        "按变化后建设用地图斑类型汇总，本期涉及建设用地图斑"
        "{{construction_change.n}}个，面积{{construction_change.area|mu}}亩。"
    )
    document.add_paragraph(
        "按单个图斑面积小于100平方米识别的细碎图斑共{{fragment_stats.n}}个，"
        "合计{{fragment_stats.area|mu}}亩。该阈值仅用于本期细碎图斑统计。"
    )

    heading("四 林地变化情况")
    document.add_paragraph(
        "林地流入{{forest_flow.inflow_m2|mu}}亩，"
        "流出{{forest_flow.outflow_m2|mu}}亩，"
        "净变化{{forest_flow.net_m2|mu}}亩；"
        "流入与流出均按变化前后地类编码比较。"
    )

    heading("五 待核实事项")
    document.add_paragraph("套合最新审批信息后，本期疑似新增违法建设用地面积***亩。")
    document.add_paragraph("其中疑似新增违法占用耕地面积***亩。")
    document.add_paragraph(
        "上述疑似违法事项需取得审批与执法信息后核实，不能以新增建设用地面积替代。"
        "面积由平方米换算为亩，保留两位小数；数字不使用千分位逗号。"
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


if __name__ == "__main__":
    print(build(Path(__file__).resolve().parents[1] / "work" / "非白名单问题库验收模板.docx"))
