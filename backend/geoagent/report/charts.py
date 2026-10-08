"""将同一批统计结果嵌入原生 Word 条形图，不依赖截图或额外绘图库。"""
from __future__ import annotations

from typing import Any
from xml.etree import ElementTree as ET

from docx.oxml import parse_xml
from docx.opc.part import Part
from docx.opc.constants import RELATIONSHIP_TYPE as RT

C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def _attach_chart(doc: Any, root: ET.Element, title: str) -> tuple[Any, Any]:
    """把图表 XML 注册到 DOCX，并返回图形段落和图题段落。"""
    package = doc.part.package
    part = Part(package.next_partname('/word/charts/chart%d.xml'),
                'application/vnd.openxmlformats-officedocument.drawingml.chart+xml',
                ET.tostring(root, encoding='utf-8', xml_declaration=True), package)
    rid = doc.part.relate_to(part, RT.CHART)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    drawing_id = doc.part.next_id
    drawing = parse_xml(f'''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:c="{C}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="5400000" cy="3200000"/>
      <wp:docPr id="{drawing_id}" name="统计图{drawing_id}"/>
      <wp:cNvGraphicFramePr/><a:graphic><a:graphicData uri="{C}"><c:chart r:id="{rid}"/>
      </a:graphicData></a:graphic></wp:inline></w:drawing>''')
    paragraph.add_run()._r.append(drawing)
    caption = doc.add_paragraph(title + '（单位：亩）')
    return paragraph, caption


def add_bar_chart(
    doc: Any,
    title: str,
    labels: list[str],
    values: list[float],
    *,
    fill_color: str = '4472C4',
    line_color: str = '2F5597',
) -> tuple[Any, Any] | None:
    """嵌入带数据缓存的二维条形图。输入值为亩；空数据不生成图。"""
    if not labels:
        return None
    if len(labels) != len(values):
        raise ValueError('图表标签和值数量不一致。')
    root = ET.Element(f'{{{C}}}chartSpace')

    def element(parent: Any, name: str, **attrs: str) -> Any:
        return ET.SubElement(parent, f'{{{C}}}{name}', attrs)

    chart = element(root, 'chart')
    element(chart, 'autoTitleDeleted', val='1')
    plot = element(chart, 'plotArea')
    element(plot, 'layout')
    bar = element(plot, 'barChart')
    element(bar, 'barDir', val='bar')
    element(bar, 'grouping', val='clustered')
    series = element(bar, 'ser')
    element(series, 'idx', val='0')
    element(series, 'order', val='0')
    element(element(series, 'tx'), 'v').text = '面积（亩）'
    # Word 会对负值柱体默认启用“反转颜色”，在部分主题和 PDF 导出时表现为白色。
    # 显式关闭反转并固定系列填充，保证正负值在 Word/PDF 中均可见。
    element(series, 'invertIfNegative', val='0')
    shape = element(series, 'spPr')
    solid_fill = ET.SubElement(shape, f'{{{A}}}solidFill')
    ET.SubElement(solid_fill, f'{{{A}}}srgbClr', {'val': fill_color})
    line = ET.SubElement(shape, f'{{{A}}}ln')
    line_fill = ET.SubElement(line, f'{{{A}}}solidFill')
    ET.SubElement(line_fill, f'{{{A}}}srgbClr', {'val': line_color})
    categories = element(element(series, 'cat'), 'strLit')
    element(categories, 'ptCount', val=str(len(labels)))
    for index, label in enumerate(labels):
        element(element(categories, 'pt', idx=str(index)), 'v').text = label
    numbers = element(element(series, 'val'), 'numLit')
    element(numbers, 'formatCode').text = '0.00'
    element(numbers, 'ptCount', val=str(len(values)))
    for index, value in enumerate(values):
        element(element(numbers, 'pt', idx=str(index)), 'v').text = f'{value:.2f}'
    data_labels = element(bar, 'dLbls')
    element(data_labels, 'numFmt', formatCode='0.00', sourceLinked='0')
    element(data_labels, 'showLegendKey', val='0')
    element(data_labels, 'showVal', val='1')
    element(data_labels, 'showCatName', val='0')
    element(data_labels, 'showSerName', val='0')
    element(bar, 'gapWidth', val='80')
    element(bar, 'axId', val='100')
    element(bar, 'axId', val='200')
    for axis_name, axis_id, cross_id, position in [('catAx', '100', '200', 'l'), ('valAx', '200', '100', 'b')]:
        axis = element(plot, axis_name)
        element(axis, 'axId', val=axis_id)
        element(element(axis, 'scaling'), 'orientation', val='minMax')
        element(axis, 'delete', val='0')
        element(axis, 'axPos', val=position)
        if axis_name == 'valAx':
            element(axis, 'majorGridlines')
            element(axis, 'numFmt', formatCode='0.00', sourceLinked='0')
        element(axis, 'tickLblPos', val='nextTo')
        element(axis, 'crossAx', val=cross_id)
        element(axis, 'crosses', val='autoZero')
        if axis_name == 'catAx':
            element(axis, 'auto', val='1')
            element(axis, 'lblAlgn', val='ctr')
            element(axis, 'lblOffset', val='100')
        else:
            element(axis, 'crossBetween', val='between')
    element(chart, 'plotVisOnly', val='1')
    element(chart, 'dispBlanksAs', val='gap')
    return _attach_chart(doc, root, title)


def add_pie_chart(
    doc: Any,
    title: str,
    labels: list[str],
    values: list[float],
    *,
    colors: list[str] | None = None,
) -> tuple[Any, Any] | None:
    """嵌入带分类名称和百分比标签的二维饼图。输入值为亩。"""
    if not labels:
        return None
    if len(labels) != len(values):
        raise ValueError('图表标签和值数量不一致。')
    palette = colors or [
        '4472C4', 'ED7D31', 'A5A5A5', 'FFC000', '5B9BD5',
        '70AD47', '264478', '9E480E', '636363', '997300',
    ]
    root = ET.Element(f'{{{C}}}chartSpace')

    def element(parent: Any, name: str, **attrs: str) -> Any:
        return ET.SubElement(parent, f'{{{C}}}{name}', attrs)

    chart = element(root, 'chart')
    element(chart, 'autoTitleDeleted', val='1')
    plot = element(chart, 'plotArea')
    element(plot, 'layout')
    pie = element(plot, 'pieChart')
    element(pie, 'varyColors', val='1')
    series = element(pie, 'ser')
    element(series, 'idx', val='0')
    element(series, 'order', val='0')
    element(element(series, 'tx'), 'v').text = '面积（亩）'
    categories = element(element(series, 'cat'), 'strLit')
    element(categories, 'ptCount', val=str(len(labels)))
    numbers = element(element(series, 'val'), 'numLit')
    element(numbers, 'formatCode').text = '0.00'
    element(numbers, 'ptCount', val=str(len(values)))
    for index, (label, value) in enumerate(zip(labels, values)):
        element(element(categories, 'pt', idx=str(index)), 'v').text = label
        element(element(numbers, 'pt', idx=str(index)), 'v').text = f'{value:.2f}'
        point = element(series, 'dPt')
        element(point, 'idx', val=str(index))
        shape = element(point, 'spPr')
        solid_fill = ET.SubElement(shape, f'{{{A}}}solidFill')
        ET.SubElement(solid_fill, f'{{{A}}}srgbClr', {'val': palette[index % len(palette)]})
    labels_node = element(pie, 'dLbls')
    element(labels_node, 'showLegendKey', val='0')
    element(labels_node, 'showVal', val='0')
    element(labels_node, 'showCatName', val='1')
    element(labels_node, 'showSerName', val='0')
    element(labels_node, 'showPercent', val='1')
    element(labels_node, 'showLeaderLines', val='1')
    legend = element(chart, 'legend')
    element(legend, 'legendPos', val='r')
    element(legend, 'overlay', val='0')
    element(chart, 'plotVisOnly', val='1')
    element(chart, 'dispBlanksAs', val='gap')
    return _attach_chart(doc, root, title)
