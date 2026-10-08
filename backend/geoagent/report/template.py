"""本地 DOCX 结构解析：不执行文档指令，不调用模型或 MinerU。"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from io import BytesIO
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
V = "{urn:schemas-microsoft-com:vml}"
MAX_UPLOAD = 10 * 1024 * 1024
PLACEHOLDER = re.compile(r"\*+|＊+|_{2,}|＿{2,}|\{\{[^{}]+\}\}")
PRESENTATION_REFERENCE = re.compile(
    r"[，,]?(?:(?:具体(?:情况)?|分(?:市|县)?统计情况)?详见[图表]\s*\d+|(?:具体)?如表\s*\d+所示)"
)
SUPERLATIVE_REGION = re.compile(
    r"((?:最多|最大|最高|最少|最低)[^，。；]{0,12}?(?:为|是))"
    r"([\u4e00-\u9fff]{1,8}(?:市|县|区))"
)
DYNAMIC_ILLEGAL_COUNT = re.compile(
    r"全省(?P<value>\d+)个(?=疑似新增违法(?:建设用地|占用耕地)的县（市、区）)"
)
NUMBERED_HEADING = re.compile(
    r"^\s*(?:第[一二三四五六七八九十百0-9]+[章节]|[一二三四五六七八九十]+、|[0-9]+[.、])"
)
FIGURE_CAPTION = re.compile(r"^\s*图\s*([0-9一二三四五六七八九十]+)\s*(.*)$")
REPORT_YEAR = 2026


def _current_business_text(text: str) -> str:
    """保留原文用于追溯，只把面向用户的问数文字切换到当前业务期。"""
    return (text.replace("2024年上半年", f"{REPORT_YEAR}年第一期")
            .replace("2024年", f"{REPORT_YEAR}年").replace("上半年", "本期"))


def _safe_docx(content: bytes, filename: str) -> list[str]:
    if not filename.lower().endswith(".docx"):
        raise ValueError("首期仅支持 .docx，请将旧 .doc 另存为 .docx。")
    if len(content) > MAX_UPLOAD:
        raise ValueError("模板不得超过10MB。")
    try:
        with ZipFile(BytesIO(content)) as package:
            infos = package.infolist()
            if len(infos) > 2000 or sum(item.file_size for item in infos) > 40 * 1024 * 1024:
                raise ValueError("模板解压体积过大。")
            if any(item.flag_bits & 1 for item in infos):
                raise ValueError("不支持加密模板。")
            if any("vbaproject" in item.filename.lower() for item in infos):
                raise ValueError("不支持含宏的模板。")
            xml = package.read("word/document.xml")
            if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                raise ValueError("模板包含不支持的XML声明。")
            etree.fromstring(xml, etree.XMLParser(resolve_entities=False, no_network=True))
            return [item.filename for item in infos]
    except (BadZipFile, KeyError, etree.XMLSyntaxError) as exc:
        raise ValueError("文件不是有效的DOCX模板。") from exc


def _font_name(run: Any) -> str | None:
    rpr = run._element.rPr
    fonts = rpr.find(W + "rFonts") if rpr is not None else None
    if fonts is None:
        return run.font.name
    return fonts.get(W + "eastAsia") or fonts.get(W + "ascii") or run.font.name


def _paragraph_style(paragraph: Paragraph) -> dict[str, Any]:
    style = paragraph.style
    heading_level = None
    if style and style.name:
        matched = re.search(r"(?:Heading|标题)\s*(\d+)", style.name, re.IGNORECASE)
        if matched:
            heading_level = int(matched.group(1))
    fmt = paragraph.paragraph_format
    points = lambda value: round(value.pt, 2) if value is not None else None
    return {
        "style_id": style.style_id if style else None,
        "style_name": style.name if style else None,
        "heading_level": heading_level,
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "left_indent_pt": points(fmt.left_indent),
        "right_indent_pt": points(fmt.right_indent),
        "first_line_indent_pt": points(fmt.first_line_indent),
        "space_before_pt": points(fmt.space_before),
        "space_after_pt": points(fmt.space_after),
        "line_spacing": points(fmt.line_spacing) if hasattr(fmt.line_spacing, "pt") else fmt.line_spacing,
        "line_spacing_rule": int(fmt.line_spacing_rule) if fmt.line_spacing_rule is not None else None,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
    }


def _runs(paragraph: Paragraph) -> list[dict[str, Any]]:
    result = []
    offset = 0
    for index, run in enumerate(paragraph.runs):
        text = run.text
        result.append({
            "index": index, "text": text, "start": offset, "end": offset + len(text),
            "bold": run.bold, "italic": run.italic, "font_name": _font_name(run),
            "font_size_pt": run.font.size.pt if run.font.size else None,
            "underline": run.underline,
            "color_rgb": str(run.font.color.rgb) if run.font.color.rgb else None,
            "highlight": int(run.font.highlight_color) if run.font.highlight_color is not None else None,
        })
        offset += len(text)
    return result


def _images(paragraph: Paragraph, block_id: str) -> list[dict[str, Any]]:
    images = []
    for index, blip in enumerate(paragraph._element.iter(A + "blip")):
        relationship_id = blip.get(R + "embed")
        relationship = paragraph.part.rels.get(relationship_id) if relationship_id else None
        drawing = next((node for node in blip.iterancestors()
                        if node.tag in {WP + "inline", WP + "anchor"}), None)
        extent = drawing.find(WP + "extent") if drawing is not None else None
        images.append({
            "id": f"{block_id}:image{index}", "relationship_id": relationship_id,
            "filename": str(getattr(getattr(relationship, "target_part", None), "partname", "")) or None,
            "placement": "inline" if drawing is not None and drawing.tag == WP + "inline" else "anchor",
            "width_emu": int(extent.get("cx")) if extent is not None else None,
            "height_emu": int(extent.get("cy")) if extent is not None else None,
            "locator": {"part": "word/document.xml", "block_id": block_id, "image_index": index},
        })
    offset = len(images)
    for index, image_data in enumerate(paragraph._element.iter(V + "imagedata"), start=offset):
        relationship_id = image_data.get(R + "id")
        relationship = paragraph.part.rels.get(relationship_id) if relationship_id else None
        shape = next((node for node in image_data.iterancestors() if node.tag == V + "shape"), None)
        images.append({
            "id": f"{block_id}:image{index}", "relationship_id": relationship_id,
            "filename": str(getattr(getattr(relationship, "target_part", None), "partname", "")) or None,
            "placement": "vml", "shape_style": shape.get("style") if shape is not None else None,
            "width_emu": None, "height_emu": None,
            "locator": {"part": "word/document.xml", "block_id": block_id, "image_index": index},
        })
    return images


def _cell_info(cell: _Cell, block_id: str, row: int, column: int) -> dict[str, Any]:
    tc_pr = cell._tc.tcPr
    span = tc_pr.find(W + "gridSpan") if tc_pr is not None else None
    merge = tc_pr.find(W + "vMerge") if tc_pr is not None else None
    shading = tc_pr.find(W + "shd") if tc_pr is not None else None
    borders = tc_pr.find(W + "tcBorders") if tc_pr is not None else None
    margins = tc_pr.find(W + "tcMar") if tc_pr is not None else None
    return {
        "id": f"{block_id}:r{row}:c{column}", "text": cell.text, "column": column,
        "grid_span": int(span.get(W + "val", "1")) if span is not None else 1,
        "vertical_merge": merge.get(W + "val", "continue") if merge is not None else None,
        "width_twips": round(cell.width.twips) if cell.width is not None else None,
        "vertical_alignment": int(cell.vertical_alignment) if cell.vertical_alignment is not None else None,
        "shading_fill": shading.get(W + "fill") if shading is not None else None,
        "borders": {
            edge.tag.removeprefix(W): {key.removeprefix(W): value for key, value in edge.attrib.items()}
            for edge in borders
        } if borders is not None else {},
        "margins_twips": {
            edge.tag.removeprefix(W): edge.get(W + "w") for edge in margins
        } if margins is not None else {},
        "paragraphs": [{"text": paragraph.text, "style": _paragraph_style(paragraph),
                        "runs": _runs(paragraph)} for paragraph in cell.paragraphs],
        "locator": {"part": "word/document.xml", "block_id": block_id, "row": row, "cell": column},
    }


def _data_start(rows: list[list[dict[str, Any]]]) -> int:
    """识别表头结束位置，避免把表头中的空白格作为回填目标。"""
    if len(rows) < 2:
        return 1
    column_count = max((sum(cell["grid_span"] for cell in row) for row in rows), default=1)
    for index, row in enumerate(rows[1:], start=1):
        non_empty = sum(bool(cell["text"].strip()) for cell in row)
        if row and row[0]["text"].strip() and non_empty <= max(2, column_count // 2):
            return index
    return 1


def _column_headers(rows: list[list[dict[str, Any]]], header_rows: int) -> list[dict[str, Any]]:
    """按 gridSpan/vMerge 把多级表头重建为每一逻辑列的层级路径。"""
    column_count = max(
        (max((cell["column"] + cell["grid_span"] for cell in row), default=0) for row in rows),
        default=0,
    )
    headers = []
    for column in range(column_count):
        levels = []
        for row in rows[:header_rows]:
            cell = next(
                (candidate for candidate in row
                 if candidate["column"] <= column < candidate["column"] + candidate["grid_span"]),
                None,
            )
            if cell is None:
                continue
            text = re.sub(r"\s+", "", cell["text"])
            if text and (not levels or levels[-1] != text):
                levels.append(text)
        headers.append({"column": column, "levels": levels, "label": " / ".join(levels)})
    duplicate_labels = {
        header["label"] for header in headers
        if header["label"] and sum(item["label"] == header["label"] for item in headers) > 1
    }
    for header in headers:
        if header["label"] not in duplicate_labels:
            continue
        values = []
        for row in rows[header_rows:header_rows + 4]:
            cell = next(
                (candidate for candidate in row
                 if candidate["column"] <= header["column"]
                 < candidate["column"] + candidate["grid_span"]),
                None,
            )
            if cell and cell["text"].strip():
                values.append(re.sub(r"\s+", "", cell["text"]))
        if any(value in {"面积", "占比", "数量"} for value in values):
            header["levels"] = header["levels"][:-1] + ["统计项"]
            header["label"] = " / ".join(header["levels"])
    land_categories = {
        "湿地", "耕地", "种植园用地", "园地", "林地", "草地", "水面", "其他土地",
        "交通运输用地", "水域及水利设施用地", "建设用地", "城镇村及工矿用地",
    }
    for header in headers:
        if len(header["levels"]) > 1 and header["levels"][0] == "其中":
            if header["levels"][1] in land_categories:
                header["levels"][0] = "地类构成"
                header["label"] = " / ".join(header["levels"])
    return headers


def _question_text(source_text: str) -> tuple[str, list[str]]:
    """清理排版引用和历史排名示例，避免它们进入数据库查询条件。"""
    cleaned = PRESENTATION_REFERENCE.sub("", source_text)
    example_regions = []

    def replace_region(match: re.Match[str]) -> str:
        example_regions.append(match.group(2))
        return match.group(1) + "[待查询地区]"

    cleaned = SUPERLATIVE_REGION.sub(replace_region, cleaned)
    cleaned = re.sub(r"[，,]([。；;])", r"\1", cleaned)
    return cleaned, example_regions


def _detect_sections(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """结合 Word 标题样式和中文编号识别章节，并把章节ID挂到后续内容块。"""
    sections = []
    current_section = None
    for block in blocks:
        if block["type"] == "paragraph":
            text = block["text"].strip()
            style_level = block["style"].get("heading_level")
            if text and (style_level is not None or NUMBERED_HEADING.match(text)):
                current_section = f"section-{len(sections) + 1}"
                level = style_level or 1
                sections.append({
                    "id": current_section, "title": text, "level": level,
                    "heading_block": block["id"], "locator": block["locator"],
                })
                block["role"] = "heading"
        block["section_id"] = current_section
    return sections


def _detect_chart_anchors(blocks: list[dict[str, Any]], images: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """以图题为主锚点，并关联其前方最近的模板图片。"""
    image_by_block = {image["locator"]["block_id"]: image for image in images}
    anchors = []
    for index, block in enumerate(blocks):
        if block["type"] != "paragraph":
            continue
        matched = FIGURE_CAPTION.match(block["text"])
        if not matched:
            continue
        existing_image = None
        for previous in reversed(blocks[:index]):
            existing_image = image_by_block.get(previous["id"])
            if existing_image or (previous["type"] == "paragraph" and previous["text"].strip()):
                break
        caption = block["text"].strip()
        anchors.append({
            "id": f"chart-anchor-{len(anchors) + 1}",
            "figure_number": matched.group(1), "caption": caption,
            "caption_block": block["id"], "section_id": block.get("section_id"),
            "placement": "replace_existing_image" if existing_image else "insert_before_caption",
            "target_image_id": existing_image["id"] if existing_image else None,
            "relationship_id": existing_image.get("relationship_id") if existing_image else None,
            "chart_type_hint": "pie" if any(word in caption for word in ("构成", "占比")) else "bar",
            "locator": existing_image["locator"] if existing_image else block["locator"],
            "status": "待确认",
        })
    return anchors


def _placeholder_unit(text: str, end: int) -> tuple[str | None, str]:
    suffix = text[end:end + 6].lstrip()
    for unit in ("万亩", "平方米", "公顷", "平方公里", "亩", "%", "个"):
        if suffix.startswith(unit):
            return unit, "integer" if unit == "个" else "number"
    if suffix.startswith(("市", "县", "区")):
        return suffix[0], "region_name"
    return None, "number"


def _source_fragment(text: str, start: int, end: int) -> str:
    left = max(text.rfind(mark, 0, start) for mark in "。；;，,、") + 1
    right_candidates = [position for mark in "。；;，,、" if (position := text.find(mark, end)) >= 0]
    right = min(right_candidates) if right_candidates else len(text)
    return text[left:right].strip()


def _metric_label(fragment: str, unit: str | None, full_text: str, start: int) -> str:
    prefix = re.split(r"[。；;，,、]|以及|和", full_text[:start])[-1]
    label = _current_business_text(prefix)
    label = PLACEHOLDER.sub("", label)
    label = re.sub(r"^2026年第一期", "", label)
    label = re.sub(r"^(?:其中|根据监测情况|本期|全省)+", "", label).strip(" ，,；;。")
    if label == "面积":
        previous = re.split(r"[；;。]", full_text[:start])[-1]
        previous = PLACEHOLDER.sub("", previous)
        previous = re.sub(r"\s*个\s*$", "", previous).strip(" ，,、；;。")
        label = previous + "面积"
    if "耕地" in full_text:
        if label == "流出":
            return "耕地流出面积"
        if label in {"净增加", "净减少", "净变化"}:
            return f"耕地{label}面积"
        if "主要流向" in label:
            return "耕地流出至" + label.split("主要流向", 1)[1] + "面积"
        flow_out = full_text[:start].rfind("耕地流出中")
        flow_in = full_text[:start].rfind("耕地流入中")
        if flow_out > flow_in and label in {"农村道路", "建设用地"}:
            return f"耕地流出至{label}面积"
        if ("恢复性地类" in label or "恢复性地类" in full_text[max(0, start - 50):start]) and "流入" in label:
            return "恢复性地类流入耕地面积"
        if flow_in > flow_out and "来源于" in full_text[:start]:
            return f"耕地流入来源于{label.removeprefix('来源于')}面积"
    if unit in {"万亩", "亩", "平方米", "公顷", "平方公里"} and not label.endswith("面积"):
        label += "面积"
    return label or "待确认指标"


def _count_metric(full_text: str, start: int, fallback: str) -> str:
    """为含括号的行政区数量恢复完整业务名称，避免只留下“区）数量”。"""
    context = full_text[:start]
    rules = [
        (r"耕地净减少的县（市、区）有\s*$", "耕地净减少县级行政区数量"),
        (r"耕地净减少的地市有\s*$", "耕地净减少设区市数量"),
        (r"新增违法面积大于1000亩的县（市、区）有\s*$", "疑似新增违法建设用地超1000亩县级行政区数量"),
        (r"新增违法占耕面积大于100亩的县（市、区）有\s*$", "疑似新增违法占耕超100亩县级行政区数量"),
    ]
    for pattern, metric in rules:
        if re.search(pattern, context):
            return metric
    return fallback.removesuffix("有") + "数量"


def _table_metric_name(block: dict[str, Any]) -> str:
    title = block.get("table_title") or ""
    title = re.sub(r"^(?:附?表)\s*[0-9一二三四五六七八九十]*\s*", "", title).strip()
    if title:
        return title
    labels = "；".join(header["label"] for header in block["column_headers"])
    if "耕地流入" in labels and "耕地流出" in labels:
        return "各行政区耕地流入流出汇总表"
    if "疑似新增违法" in labels:
        return "各行政区疑似新增违法建设用地汇总表"
    if "指标" in labels and "图斑数" in labels and "面积" in labels:
        if "耕地" in block.get("caption", ""):
            return "耕地变化总体情况表"
        return "总体指标汇总表"
    return "指标汇总表"


def decompose_atomic_items(parsed: dict[str, Any]) -> list[dict[str, Any]]:
    """把每个段落占位符或表格区域转成可匹配、可回填的原子查询项。"""
    block_by_id = {block["id"]: block for block in parsed["blocks"]}
    items = []
    region_counts: dict[str, int] = {}
    last_region_by_block: dict[str, dict[str, Any]] = {}
    last_count_by_block: dict[str, dict[str, Any]] = {}
    block_order = {block["id"]: index for index, block in enumerate(parsed["blocks"])}
    ordered_slots = sorted(
        parsed["slots"],
        key=lambda slot: (
            block_order[slot["block_id"]],
            slot.get("start", 10**9),
            slot["id"],
        ),
    )
    for slot in ordered_slots:
        block = block_by_id[slot["block_id"]]
        item_id = f"field-{len(items) + 1:03d}"
        match_stub = {
            "status": "unmatched", "matched_question_id": None, "confidence": None,
            "match_reasons": [], "candidates": [],
        }
        if slot["kind"] in {"placeholder", "dynamic_region", "dynamic_number"}:
            source_text = block["text"]
            if slot["kind"] == "dynamic_region":
                unit, expected_type = "地区名称", "region_name"
            elif slot["kind"] == "dynamic_number":
                unit, expected_type = "个", "integer"
            else:
                unit, expected_type = _placeholder_unit(source_text, slot["end"])
            fragment = _source_fragment(source_text, slot["start"], slot["end"])
            metric = _metric_label(fragment, unit, source_text, slot["start"])
            related_to = None
            if expected_type == "region_name":
                region_counts[block["id"]] = region_counts.get(block["id"], 0) + 1
                number = region_counts[block["id"]]
                context = source_text[:slot["start"]]
                if "违法占耕最多" in fragment or "违法占耕最多" in context[-40:]:
                    metric = "疑似新增违法占耕最多地区"
                elif "新增违法面积大于1000亩" in context:
                    metric = f"疑似新增违法建设用地超1000亩排名地区{number}"
                elif "耕地净减少" in source_text:
                    context = source_text[:slot["start"]]
                    metric = (
                        "耕地净减少最多县级行政区"
                        if context.rfind("县（市、区）") > context.rfind("地市")
                        else "耕地净减少最多设区市"
                    )
                else:
                    metric = f"动态排名地区{number}"
                last_region_by_block[block["id"]] = {"id": item_id, "metric": metric}
            elif unit in {"万亩", "亩"} and "为" in fragment and block["id"] in last_region_by_block:
                region_item = last_region_by_block[block["id"]]
                metric = region_item["metric"] + "面积"
                related_to = region_item["id"]
            elif unit == "个":
                metric = _count_metric(source_text, slot["start"], metric)
                nearby = source_text[max(0, slot["start"] - 90):slot["start"]]
                if "恢复性地类" in nearby and "转为耕地" in nearby:
                    metric = "恢复性地类流入耕地图斑数量"
            elif unit in {"万亩", "亩", "平方米", "公顷", "平方公里"} and block["id"] in last_count_by_block:
                metric = last_count_by_block[block["id"]]["metric"].removesuffix("数量") + "面积"
            elif unit == "%" and metric.endswith("面积的"):
                metric = metric.removesuffix("的") + "占比"
            context = source_text[:slot["start"]]
            if metric == "耕地净减少面积":
                if context.rfind("县（市、区）") > context.rfind("地市"):
                    metric = "耕地净减少最多县级行政区面积"
                elif "地市" in context:
                    metric = "耕地净减少最多设区市面积"
            _, historical_examples = _question_text(_current_business_text(fragment))
            item = {
                "id": item_id, "kind": "scalar", "section_id": block.get("section_id"),
                "source_block": block["id"], "source_text": source_text,
                "source_fragment": fragment, "metric": metric, "scope": "全省" if "全省" in source_text else None,
                "period": f"{REPORT_YEAR}年第一期", "unit": unit, "expected_type": expected_type,
                "query_question": f"查询{REPORT_YEAR}年第一期{metric}",
                "historical_examples": historical_examples,
                "binding": {"target_type": (
                                "paragraph_placeholder" if slot["kind"] == "placeholder"
                                else "paragraph_dynamic_text"
                            ), "slot_id": slot["id"],
                            "locator": slot["locator"], "format": {"unit": unit, "thousands_separator": False}},
                "question_library_match": match_stub, "status": "待用户确认",
            }
            if related_to:
                item["related_to"] = related_to
            items.append(item)
            if unit == "个":
                last_count_by_block[block["id"]] = {"id": item_id, "metric": metric}
        else:
            headers = [header["label"] for header in block["column_headers"] if header["label"]]
            items.append({
                "id": item_id, "kind": "table", "section_id": block.get("section_id"),
                "source_block": block["id"], "source_text": block.get("table_title") or block["caption"],
                "metric": _table_metric_name(block), "scope": "全省",
                "period": f"{REPORT_YEAR}年第一期", "unit": block.get("unit"), "expected_type": "table",
                "query_question": "按行政区划汇总：" + "；".join(headers),
                "result_schema": headers,
                "binding": {"target_type": "table_region", "slot_id": slot["id"],
                            "target_cell_ids": slot["target_cell_ids"], "locator": slot["locator"]},
                "question_library_match": match_stub, "status": "待用户确认",
            })
    return items


def _iter_blocks(document: DocumentObject) -> list[Paragraph | Table]:
    if hasattr(document, "iter_inner_content"):
        return list(document.iter_inner_content())
    blocks: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        if child.tag == W + "p":
            blocks.append(Paragraph(child, document))
        elif child.tag == W + "tbl":
            blocks.append(Table(child, document))
    return blocks


def parse_docx(content: bytes, filename: str) -> dict[str, Any]:
    """用 python-docx 读取常规结构，并用底层 OOXML 补充定位和合并信息。"""
    package_names = _safe_docx(content, filename)
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("文件不是有效的DOCX模板。") from exc

    warnings = ["JSON中的文字来自用户模板，不作为系统指令执行。"]
    if any(re.fullmatch(r"word/(header|footer)\d+\.xml", name) for name in package_names):
        warnings.append("模板含页眉或页脚，已记录部件名称；其中日期和单位仍需人工复核。")

    blocks: list[dict[str, Any]] = []
    slots: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    previous_text = ""
    recent_paragraphs: list[dict[str, str]] = []
    all_items = _iter_blocks(document)
    for index, item in enumerate(all_items):
        block_id = f"b{index}"
        if isinstance(item, Paragraph):
            text = item.text
            paragraph_images = _images(item, block_id)
            images.extend(paragraph_images)
            blocks.append({
                "id": block_id, "type": "paragraph", "text": text,
                "style": _paragraph_style(item), "runs": _runs(item), "images": paragraph_images,
                "locator": {"part": "word/document.xml", "block_index": index},
            })
            for number, match in enumerate(PLACEHOLDER.finditer(text)):
                slots.append({
                    "id": f"{block_id}:s{number}", "block_id": block_id,
                    "kind": "placeholder", "text": match.group(), "start": match.start(),
                    "end": match.end(), "context": text,
                    "locator": {"part": "word/document.xml", "block_id": block_id,
                                "start": match.start(), "end": match.end()},
                })
            occupied = [(slot["start"], slot["end"]) for slot in slots if slot["block_id"] == block_id]
            for number, match in enumerate(SUPERLATIVE_REGION.finditer(text)):
                start, end = match.span(2)
                if any(left < end and right > start for left, right in occupied):
                    continue
                slots.append({
                    "id": f"{block_id}:region{number}", "block_id": block_id,
                    "kind": "dynamic_region", "text": match.group(2),
                    "start": start, "end": end, "context": text,
                    "locator": {"part": "word/document.xml", "block_id": block_id,
                                "start": start, "end": end},
                })
                occupied.append((start, end))
            for number, match in enumerate(DYNAMIC_ILLEGAL_COUNT.finditer(text)):
                start, end = match.span("value")
                slots.append({
                    "id": f"{block_id}:number{number}", "block_id": block_id,
                    "kind": "dynamic_number", "text": match.group("value"),
                    "start": start, "end": end, "context": text,
                    "locator": {"part": "word/document.xml", "block_id": block_id,
                                "start": start, "end": end},
                })
            if text.strip():
                previous_text = text.strip()
                recent_paragraphs.append({"id": block_id, "text": text.strip()})
                recent_paragraphs = recent_paragraphs[-5:]
            continue

        rows = []
        for ri, row in enumerate(item.rows):
            cells = []
            column = 0
            for tc in row._tr.tc_lst:
                cell = _cell_info(_Cell(tc, item), block_id, ri, column)
                cells.append(cell)
                column += cell["grid_span"]
            rows.append(cells)
        start = _data_start(rows)
        column_headers = _column_headers(rows, start)
        fill_targets: list[str] = []
        explicit_placeholders: list[str] = []
        for ri, row in enumerate(rows):
            for cell in row:
                continuation = cell["vertical_merge"] == "continue"
                matches = list(PLACEHOLDER.finditer(cell["text"]))
                if matches and not continuation:
                    explicit_placeholders.append(cell["id"])
                if ri >= start and not continuation and (
                    matches or (cell["column"] > 0 and not cell["text"].strip())
                ):
                    fill_targets.append(cell["id"])
        unit_text = next((item["text"] for item in reversed(recent_paragraphs)
                          if re.search(r"单位\s*[：:]", item["text"])), "")
        unit_match = re.search(r"单位\s*[：:]\s*([^\s；;，,。]+)", unit_text)
        non_unit = [item["text"] for item in reversed(recent_paragraphs)
                    if not re.search(r"单位\s*[：:]", item["text"])]
        table_title = ""
        if non_unit:
            candidate = non_unit[0]
            if re.match(r"^(?:附?表)\s*[0-9一二三四五六七八九十]+", candidate):
                table_title = candidate
            elif len(non_unit) > 1 and re.match(
                r"^附表\s*[0-9一二三四五六七八九十]+", non_unit[1]
            ):
                table_title = candidate
        blocks.append({
            "id": block_id, "type": "table", "caption": previous_text,
            "table_title": table_title, "unit": unit_match.group(1) if unit_match else None,
            "context_paragraphs": list(recent_paragraphs),
            "style": {
                "style_name": item.style.name if item.style else None,
                "alignment": int(item.alignment) if item.alignment is not None else None,
                "autofit": item.autofit,
                "grid_widths_twips": [
                    int(grid_col.get(W + "w")) if grid_col.get(W + "w") else None
                    for grid_col in item._tbl.tblGrid
                ],
                "row_heights_twips": [
                    round(row.height.twips) if row.height is not None else None
                    for row in item.rows
                ],
            },
            "header_rows": start, "column_headers": column_headers,
            "rows": rows, "fill_targets": fill_targets,
            "locator": {"part": "word/document.xml", "block_index": index},
        })
        if fill_targets or explicit_placeholders:
            slots.append({
                "id": f"{block_id}:table", "block_id": block_id, "kind": "table_region",
                "context": previous_text, "target_cell_ids": fill_targets,
                "explicit_placeholder_cell_ids": explicit_placeholders,
                "locator": {"part": "word/document.xml", "block_id": block_id,
                            "data_start_row": start},
            })

    if any(item._element.find(".//" + W + "txbxContent") is not None for item in all_items):
        warnings.append("发现文本框，首期仅记录提示，不自动回填文本框。")
    sections = _detect_sections(blocks)
    chart_anchors = _detect_chart_anchors(blocks, images)
    parsed_at = datetime.now().astimezone()
    return {
        "schema_version": 4, "filename": filename, "sha256": sha256(content).hexdigest(),
        "parsed_at": parsed_at.isoformat(timespec="seconds"), "report_year": REPORT_YEAR,
        "sections": sections, "blocks": blocks, "slots": slots, "images": images,
        "chart_anchors": chart_anchors,
        "package_parts": {
            "headers": [name for name in package_names if re.fullmatch(r"word/header\d+\.xml", name)],
            "footers": [name for name in package_names if re.fullmatch(r"word/footer\d+\.xml", name)],
            "media": [name for name in package_names if name.startswith("word/media/") and not name.endswith("/")],
        },
        "warnings": warnings,
    }


def decompose_questions(parsed: dict[str, Any]) -> list[dict[str, Any]]:
    """生成带原文对照和来源定位的待审核问题，不使用预置问题列表。"""
    slots_by_block: dict[str, list[dict[str, Any]]] = {}
    for slot in parsed["slots"]:
        slots_by_block.setdefault(slot["block_id"], []).append(slot)
    atomic_items = decompose_atomic_items(parsed)
    atomic_by_block: dict[str, list[dict[str, Any]]] = {}
    for item in atomic_items:
        atomic_by_block.setdefault(item["source_block"], []).append(item)
    questions = []
    for block in parsed["blocks"]:
        block_slots = slots_by_block.get(block["id"], [])
        if not block_slots:
            continue
        if block["type"] == "paragraph":
            source_text = block["text"]
            query_text, example_regions = _question_text(_current_business_text(source_text))
            question = f"请查询这段内容需要的统计数据：{query_text}"
        else:
            source_text = _table_metric_name(block) + "；表头：" + "；".join(
                header["label"] for header in block["column_headers"] if header["label"]
            )
            query_text, example_regions = _question_text(_current_business_text(source_text))
            question = f"请按照这张表的栏目汇总数据：{query_text}"
        requirements = "面积统一为亩，数字不加千分位逗号；缺失数据标注待核实，不填零。"
        if example_regions:
            requirements += (
                " 模板中的" + "、".join(dict.fromkeys(example_regions))
                + "是历史示例，不得作为筛选条件；必须从当前数据库动态确定排名地区。"
            )
        questions.append({
            "id": f"Q{len(questions) + 1}", "source_block": block["id"],
            "section_id": block.get("section_id"),
            "source_text": source_text, "source_locator": block["locator"],
            "slot_ids": [slot["id"] for slot in block_slots], "question": question,
            "atomic_item_ids": [item["id"] for item in atomic_by_block.get(block["id"], [])],
            "requirements": requirements, "historical_examples": example_regions,
            "question_library_match": {
                "status": "unmatched", "matched_question_id": None, "confidence": None,
                "match_reasons": [], "candidates": [],
            },
            "status": "待用户确认", "requires_business_review": True,
        })
    return questions
