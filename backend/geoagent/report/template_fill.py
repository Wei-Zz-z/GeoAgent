"""依据模板解析 JSON 中的稳定定位信息回填 DOCX。"""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .charts import add_bar_chart, add_pie_chart
from .template import _iter_blocks


def _replace_range(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    """跨 run 替换文本，尽量保留占位符所在首个 run 的样式。"""
    runs = list(paragraph.runs)
    offsets: list[tuple[int, int]] = []
    cursor = 0
    for run in runs:
        offsets.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)
    touched = [index for index, (left, right) in enumerate(offsets) if left < end and right > start]
    if not touched:
        raise ValueError(f"段落定位失效：找不到字符区间 {start}:{end}")
    first, last = touched[0], touched[-1]
    first_left, first_right = offsets[first]
    last_left, _ = offsets[last]
    prefix = runs[first].text[: max(start - first_left, 0)]
    suffix = runs[last].text[max(end - last_left, 0):]
    runs[first].text = prefix + replacement + suffix
    for index in touched[1:]:
        runs[index].text = ""


def _set_cell_text(cell: Any, value: Any) -> None:
    """在原单元格段落中改文字，保留段落、文字及单元格样式。"""
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = str(value)
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(str(value))
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def _append_data_row(table: Table, data_start: int) -> None:
    """扩充数据行时继承原数据行的尺寸、边框、底纹和字体。"""
    if len(table.rows) <= data_start:
        table.add_row()
        return
    row_xml = deepcopy(table.rows[-1]._tr)
    table._tbl.append(row_xml)
    for cell in table.rows[-1].cells:
        _set_cell_text(cell, "")


def refill_docx(
    template_path: str | Path,
    parsed: dict[str, Any],
    output_path: str | Path,
    *,
    scalar_answers: dict[str, Any],
    table_answers: dict[str, list[list[Any]]] | None = None,
    table_cell_answers: dict[str, str] | None = None,
    charts: dict[str, dict[str, Any]] | None = None,
    test_notice: str | None = None,
    text_replacements: dict[str, str] | None = None,
    remove_chart_anchors: bool = False,
) -> Path:
    """按 slot id 回填文本/表格，并按 chart anchor id 插入原生 Word 图表。"""
    if sha256(Path(template_path).read_bytes()).hexdigest() != parsed["sha256"]:
        raise ValueError("模板内容与解析时不同，请重新上传并解析后再回填。")
    document = Document(str(template_path))
    blocks = _iter_blocks(document)
    slots_by_block: dict[str, list[dict[str, Any]]] = {}
    for slot in parsed["slots"]:
        slots_by_block.setdefault(slot["block_id"], []).append(slot)
    allowed_cells = {
        cell_id for slot in parsed["slots"] if slot["kind"] == "table_region"
        for cell_id in slot["target_cell_ids"]
    }
    if set(table_cell_answers or {}) - allowed_cells:
        raise ValueError("表格绑定包含模板中不存在的可填单元格。")

    for block in parsed["blocks"]:
        block_slots = slots_by_block.get(block["id"], [])
        if not block_slots:
            continue
        target = blocks[block["locator"]["block_index"]]
        paragraph_slots = [
            slot for slot in block_slots
            if slot["kind"] in {"placeholder", "dynamic_region", "dynamic_number"}
            and slot["id"] in scalar_answers
        ]
        if paragraph_slots:
            if not isinstance(target, Paragraph):
                raise ValueError(f"回填定位失效：{block['id']} 不是段落。")
            if target.text != block["text"]:
                raise ValueError(f"回填定位失效：{block['id']} 原文已变化。")
            for slot in sorted(paragraph_slots, key=lambda item: item["start"], reverse=True):
                _replace_range(target, slot["start"], slot["end"], str(scalar_answers[slot["id"]]))

        table_slot = next((slot for slot in block_slots if slot["kind"] == "table_region"), None)
        rows = (table_answers or {}).get(table_slot["id"] if table_slot else "")
        if table_slot and rows is not None:
            if not isinstance(target, Table):
                raise ValueError(f"回填定位失效：{block['id']} 不是表格。")
            start = table_slot["locator"]["data_start_row"]
            while len(target.rows) < start + len(rows):
                _append_data_row(target, start)
            for row_index, values in enumerate(rows, start=start):
                cells = target.rows[row_index].cells
                for column, value in enumerate(values[: len(cells)]):
                    _set_cell_text(cells[column], value)
            for row_index in range(start + len(rows), len(target.rows)):
                for cell in target.rows[row_index].cells:
                    _set_cell_text(cell, "")
        if table_slot and table_cell_answers:
            if not isinstance(target, Table):
                raise ValueError(f"回填定位失效：{block['id']} 不是表格。")
            for parsed_row in block["rows"]:
                for parsed_cell in parsed_row:
                    cell_id = parsed_cell["id"]
                    if cell_id not in table_cell_answers:
                        continue
                    locator = parsed_cell["locator"]
                    cell = target.cell(locator["row"], locator["cell"])
                    if cell.text != parsed_cell["text"]:
                        raise ValueError(f"回填定位失效：{cell_id} 原文已变化。")
                    _set_cell_text(cell, table_cell_answers[cell_id])

    if remove_chart_anchors:
        for anchor in parsed.get("chart_anchors", []):
            if anchor.get("target_image_id"):
                image_block_id = anchor["target_image_id"].split(":image", 1)[0]
                image_block = next(block for block in parsed["blocks"] if block["id"] == image_block_id)
                image_target = blocks[image_block["locator"]["block_index"]]
                if isinstance(image_target, Paragraph) and not image_target.text.strip():
                    image_target._element.getparent().remove(image_target._element)
            caption_block = next(
                block for block in parsed["blocks"] if block["id"] == anchor["caption_block"]
            )
            caption_target = blocks[caption_block["locator"]["block_index"]]
            if isinstance(caption_target, Paragraph):
                caption_target._element.getparent().remove(caption_target._element)

    for anchor in [] if remove_chart_anchors else parsed.get("chart_anchors", []):
        spec = (charts or {}).get(anchor["id"])
        if not spec:
            continue
        labels = [str(item) for item in spec["labels"]]
        values = [float(item) for item in spec["values"]]
        if spec.get("chart_type", anchor.get("chart_type_hint")) == "pie":
            result = add_pie_chart(
                document, anchor["caption"], labels, values,
                colors=spec.get("colors"),
            )
        else:
            result = add_bar_chart(
                document, anchor["caption"], labels, values,
                fill_color=spec.get("fill_color", "4472C4"),
                line_color=spec.get("line_color", "2F5597"),
            )
        if result is None:
            continue
        chart_paragraph, generated_caption = result
        caption_index = next(
            block["locator"]["block_index"]
            for block in parsed["blocks"]
            if block["id"] == anchor["caption_block"]
        )
        caption = blocks[caption_index]
        if not isinstance(caption, Paragraph):
            raise ValueError("图表锚点定位失效：图题不是段落。")
        if anchor.get("placement") == "replace_existing_image" and anchor.get("target_image_id"):
            image_block_id = anchor["target_image_id"].split(":image", 1)[0]
            image_block = next(block for block in parsed["blocks"] if block["id"] == image_block_id)
            image_target = blocks[image_block["locator"]["block_index"]]
            if isinstance(image_target, Paragraph):
                image_target._p.addprevious(chart_paragraph._p)
                if not image_target.text.strip():
                    image_target._element.getparent().remove(image_target._element)
            else:
                caption._p.addprevious(chart_paragraph._p)
        else:
            caption._p.addprevious(chart_paragraph._p)
        generated_caption._element.getparent().remove(generated_caption._element)

    if text_replacements:
        paragraphs = list(document.paragraphs)
        paragraphs.extend(
            paragraph
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        for paragraph in paragraphs:
            for old, new in text_replacements.items():
                start = paragraph.text.rfind(old)
                while start >= 0:
                    _replace_range(paragraph, start, start + len(old), new)
                    start = paragraph.text.rfind(old, 0, start)

    if test_notice:
        notice = document.paragraphs[0].insert_paragraph_before(test_notice)
        notice.runs[0].bold = True
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    if text_replacements:
        temporary = output.with_name(output.stem + ".rewrite.docx")
        with ZipFile(output) as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                content = source.read(info.filename)
                package_only_parts = (
                    info.filename in {"word/footnotes.xml", "word/endnotes.xml"}
                    or info.filename.startswith("word/header")
                    or info.filename.startswith("word/footer")
                )
                if package_only_parts and info.filename.endswith(".xml"):
                    for old, new in text_replacements.items():
                        content = content.replace(old.encode("utf-8"), new.encode("utf-8"))
                target.writestr(info, content)
        temporary.replace(output)
    return output
