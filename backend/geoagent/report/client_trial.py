"""甲方固定快报模板的最小自动问数与回填试运行。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..tools.categories import category_codes
from .template_fill import refill_docx


TEMPLATE_SHA256 = "b4757efdb23483738e8cb1255d31567ae3648ec349eb17a648581937fcf48c56"
TABLE = 'data."2026_1_change_landuse"'
CITY_NAMES = {
    "3301": "杭州市", "3302": "宁波市", "3303": "温州市", "3304": "嘉兴市",
    "3305": "湖州市", "3306": "绍兴市", "3307": "金华市", "3308": "衢州市",
    "3309": "舟山市", "3310": "台州市", "3311": "丽水市",
}


def is_client_trial_template(parsed: dict[str, Any]) -> bool:
    """只接受已核验的那一份 DOCX，不把相似模板误作固定口径。"""
    return (
        parsed.get("sha256") == TEMPLATE_SHA256
        and len(parsed.get("blocks", [])) == 41
        and len(parsed.get("slots", [])) == 35
        and [block.get("table_title") for block in parsed["blocks"] if block["type"] == "table"] == [
            "表1 各市耕地净变化情况表",
            "表2耕地净减少最多的10个县（市、区）",
            "表3 各市疑似新增违法建设用地情况表",
            "表4 疑似新增违法占耕面积大于100亩的县（市、区）",
            "耕地总体变化情况统计表",
            "疑似新增违法建设用地情况统计表",
        ]
    )


def _codes_in(codes: list[str]) -> str:
    return ", ".join(f"'{code}'" for code in codes)


async def query_client_trial(gateway: Any, skills_dir: Path) -> dict[str, Any]:
    """固定只读 SQL；所有数值保留平方米原值，显示阶段再换算为亩。"""
    crop = '"TBLX" IN (\'01\', \'1\')'
    orig = '"DLBM" LIKE \'01%\''
    construction_codes = category_codes("construction", skills_dir)
    construction = f'"TBLX" IN ({_codes_in(construction_codes)})'
    inflow = f"({crop} AND NOT {orig})"
    outflow = f"({orig} AND NOT {crop})"
    sql = (
        'SELECT '
        f'COALESCE(SUM(CASE WHEN {inflow} THEN "MJ" END),0) AS inflow, '
        f'COALESCE(SUM(CASE WHEN {outflow} THEN "MJ" END),0) AS outflow, '
        f'COALESCE(SUM(CASE WHEN {orig} AND {construction} THEN "MJ" END),0) AS to_construction, '
        f'COALESCE(SUM(CASE WHEN {inflow} AND "DLBM" LIKE \'02%\' THEN "MJ" END),0) AS from_orchard, '
        f'COALESCE(SUM(CASE WHEN {inflow} AND "DLBM" LIKE \'03%\' THEN "MJ" END),0) AS from_forest, '
        f'COALESCE(SUM(CASE WHEN {inflow} AND "DLBM" ~ \'^(05|06|07|08|09|100[1-5]|100[7-9]|1109|1201)\' THEN "MJ" END),0) AS from_construction, '
        f'COALESCE(SUM(CASE WHEN {inflow} AND "DLBM" IN (\'1104\',\'1104A\',\'1104K\') THEN "MJ" END),0) AS from_pond, '
        f'COALESCE(SUM(CASE WHEN {construction} THEN "MJ" END),0) AS construction '
        f'FROM {TABLE}'
    )
    total_rows = (await gateway.run_sql(sql, conversation_id="client_trial"))["rows"]
    if len(total_rows) != 1:
        raise ValueError("固定统计未返回唯一全省汇总行。")

    area_parts = (
        f'COALESCE(SUM(CASE WHEN {inflow} THEN "MJ" END),0) AS inflow, '
        f'COALESCE(SUM(CASE WHEN {outflow} THEN "MJ" END),0) AS outflow'
    )
    city_sql = (
        f'SELECT LEFT("XZQDM"::text,4) AS code, {area_parts} '
        f'FROM {TABLE} GROUP BY LEFT("XZQDM"::text,4) ORDER BY code'
    )
    county_sql = (
        f'SELECT "XZQDM" AS code, "XMC" AS name, {area_parts} '
        f'FROM {TABLE} GROUP BY "XZQDM", "XMC" ORDER BY code'
    )
    cities = (await gateway.run_sql(city_sql, conversation_id="client_trial"))["rows"]
    counties = (await gateway.run_sql(county_sql, conversation_id="client_trial"))["rows"]
    if not cities or not counties or any(str(row["code"]) not in CITY_NAMES for row in cities):
        raise ValueError("行政区代码无法对应到浙江省11个设区市，已停止自动回填。")
    return {"total": total_rows[0], "cities": cities, "counties": counties}


def _mu(area_m2: Any) -> str:
    return f"{float(area_m2) * 0.0015:.2f}"


def _net(row: dict[str, Any]) -> float:
    return float(row["inflow"]) - float(row["outflow"])


def _replace_paragraph(paragraph: Any, content: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = content
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(content)


def fill_client_trial_report(
    template_path: Path,
    parsed: dict[str, Any],
    result: dict[str, Any],
    output_path: Path,
    *,
    query_started_at: datetime,
) -> dict[str, Any]:
    """可靠指标自动回填；无管理信息的数据保留为待核实，不伪装正式报告。"""
    total = result["total"]
    cities = sorted(result["cities"], key=lambda row: list(CITY_NAMES).index(str(row["code"])))
    counties = result["counties"]
    decreasing_cities = [row for row in cities if _net(row) < 0]
    decreasing_counties = sorted((row for row in counties if _net(row) < 0), key=_net)
    worst_city = min(decreasing_cities, key=_net) if decreasing_cities else None
    worst_county = decreasing_counties[0] if decreasing_counties else None
    values: dict[str, str] = {
        "耕地流入面积": _mu(total["inflow"]),
        "耕地流出面积": _mu(total["outflow"]),
        "耕地净增加面积": _mu(float(total["inflow"]) - float(total["outflow"])),
        "耕地流出至建设用地面积": _mu(total["to_construction"]),
        "耕地流入来源于园地面积": _mu(total["from_orchard"]),
        "耕地流入来源于林地面积": _mu(total["from_forest"]),
        "耕地流入来源于建设用地面积": _mu(total["from_construction"]),
        "耕地流入来源于坑塘水面面积": _mu(total["from_pond"]),
        "耕地净减少设区市数量": str(len(decreasing_cities)),
        "耕地净减少县级行政区数量": str(len(decreasing_counties)),
        "新增建设用地面积": _mu(total["construction"]),
    }
    if worst_city:
        values["耕地净减少最多设区市"] = CITY_NAMES[str(worst_city["code"])]
        values["耕地净减少设区市面积"] = _mu(-_net(worst_city))
    if worst_county:
        values["耕地净减少最多县级行政区"] = str(worst_county["name"])
        values["耕地净减少县级行政区面积"] = _mu(-_net(worst_county))

    from .template import decompose_atomic_items

    items = decompose_atomic_items(parsed)
    scalar_answers = {
        item["binding"]["slot_id"]: values.get(item["metric"], "待核实")
        for item in items if item["kind"] == "scalar"
    }
    table_answers = {}
    for slot in parsed["slots"]:
        if slot["kind"] != "table_region":
            continue
        title = next(block for block in parsed["blocks"] if block["id"] == slot["block_id"])["table_title"]
        if title == "表1 各市耕地净变化情况表":
            table_answers[slot["id"]] = [
                ["合计", _mu(total["outflow"]), _mu(total["inflow"]), _mu(float(total["inflow"]) - float(total["outflow"]))],
                *[[CITY_NAMES[str(row["code"])], _mu(row["outflow"]), _mu(row["inflow"]), _mu(_net(row))] for row in cities],
            ]
        elif title == "表2耕地净减少最多的10个县（市、区）":
            table_answers[slot["id"]] = [
                [str(row["name"]), _mu(row["outflow"]), _mu(row["inflow"]), _mu(_net(row))]
                for row in decreasing_counties[:10]
            ]
    charts = {}
    for anchor in parsed.get("chart_anchors", []):
        if anchor["caption"] == "图1 耕地流入和流出情况":
            charts[anchor["id"]] = {
                "chart_type": "bar", "labels": ["耕地流入", "耕地流出"],
                "values": [float(total["inflow"]) * 0.0015, float(total["outflow"]) * 0.0015],
                "fill_color": "70AD47", "line_color": "548235",
            }
    refill_docx(
        template_path, parsed, output_path,
        scalar_answers=scalar_answers, table_answers=table_answers, charts=charts,
        test_notice="【自动化试运行版】主表可支持的数据已自动查询回填；恢复性地类、疑似违法及未完成的附表仍待核实。变化后建设用地按临时分类口径统计，不得作为正式业务报告。",
        text_replacements={
            "利用2024年3-5月份遥感影像与2023年度国土变更调查影像进行对比": "利用本期遥感影像与现有国土变更调查数据进行对比",
            "2024年全省上半年": "2026年第一期全省",
            "2024年上半年": "2026年第一期",
            "2024年6月24日": query_started_at.astimezone().strftime("%Y年%m月%d日"),
            "一、全省耕地数量稳定且略有增加": "一、全省耕地变化情况",
            "二、疑似新增违法占耕依然发生": "二、建设用地变化及待核实事项",
            "三、下一步工作": "三、待核实事项",
            "净增加": "净变化",
            "上半年": "本期",
            "万亩": "亩",
            "待核实亩": "待核实",
        },
    )
    document = Document(output_path)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("厅调查处组织测科院监测中心"):
            _replace_paragraph(
                paragraph,
                "本报告基于2026年第一期地类变化图斑数据，按固定统计口径自动汇总耕地变化和建设用地变化，"
                "用于快报自动化链路试运行。涉及用地管理信息套合的疑似违法数据尚未接入，以下相应位置标为待核实。",
            )
        elif paragraph.text.startswith("根据监测情况") and "疑似新增违法建设用地" in paragraph.text:
            _replace_paragraph(
                paragraph,
                f"根据监测情况，本期全省变化后建设用地图斑面积{_mu(total['construction'])}亩。"
                "疑似新增违法建设用地及占耕数据待接入用地管理信息套合结果后核实，分市情况见待核实的表3。",
            )
            for reference in paragraph._p.xpath(".//w:footnoteReference"):
                reference.getparent().remove(reference)
        elif "疑似新增违法建设用地的县" in paragraph.text and "1000亩" in paragraph.text:
            _replace_paragraph(paragraph, "疑似新增违法建设用地和占耕的县域数量、排名与面积，待接入用地管理信息套合结果后核实。")
        elif paragraph.text.startswith("我们已经将耕地变化情况"):
            _replace_paragraph(paragraph, "本报告为自动化链路试运行结果。疑似违法数据和未完成的附表待业务部门核实后补充。")
        elif paragraph.text.startswith(("一是执法局将", "二是耕保处将", "三是调查处将")):
            paragraph._element.getparent().remove(paragraph._element)
        elif paragraph.text.strip() in {"表3 各市疑似新增违法建设用地情况表", "表4 疑似新增违法占耕面积大于100亩的县（市、区）", "耕地总体变化情况统计表", "疑似新增违法建设用地情况统计表"}:
            paragraph.add_run("（待核实）")
        if paragraph.text.startswith("表2耕地净减少最多的10个县"):
            for page_break in paragraph._p.xpath(".//w:br[@w:type='page']"):
                page_break.getparent().remove(page_break)
    if document.tables:
        document.tables[0].cell(0, 0).text = "设区市"
    for table in document.tables[:2]:
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        if header_properties.find(qn("w:tblHeader")) is None:
            header_properties.append(OxmlElement("w:tblHeader"))
    if len(document.tables) >= 4:
        for index, row in enumerate(document.tables[3].rows[1:]):
            if row.cells[0].text.strip().startswith(("**", "＊")):
                row.cells[0].text = "待核实" if index == 0 else ""
    for table in document.tables[4:]:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() in {"100%", "%"}:
                    cell.text = ""
    document.save(output_path)
    filled = [item for item in items if item["kind"] == "scalar" and item["metric"] in values]
    pending = [item["metric"] for item in items if item["kind"] == "scalar" and item["metric"] not in values]
    return {"filled_scalars": len(filled), "pending_scalars": pending, "filled_tables": len(table_answers), "pending_tables": 6 - len(table_answers), "charts": len(charts)}
