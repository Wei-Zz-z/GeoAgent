"""从已生成的土地变化快报反向制作模板，并验证回填闭环。"""
from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

NUMBER = re.compile(r"-?\d[\d,]*(?:\.\d+)?")
STAT_PARAGRAPH_PREFIXES = (
    "本期通过遥感影像",
    "变化前为耕地的图斑",
    "原耕地流向建设用地",
    "变化后图斑类型为建设用地",
)


def _blocks(document: Document) -> list[Paragraph | Table]:
    return list(document.iter_inner_content())


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """保留首个 run 的字符格式，仅替换可编辑文字。"""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _number(text: str) -> float:
    return float(text.replace(",", ""))


def _numbers(text: str) -> list[float]:
    return [_number(item) for item in NUMBER.findall(text)]


def extract_stats(source: Document) -> dict[str, Any]:
    """从师兄成品快报读取一组可复现实验的标准答案。"""
    by_prefix = {
        prefix: next(p.text for p in source.paragraphs if p.text.startswith(prefix))
        for prefix in STAT_PARAGRAPH_PREFIXES
    }
    overall = _numbers(by_prefix[STAT_PARAGRAPH_PREFIXES[0]])
    crop = _numbers(by_prefix[STAT_PARAGRAPH_PREFIXES[1]])
    flow = _numbers(by_prefix[STAT_PARAGRAPH_PREFIXES[2]])
    construction = _numbers(by_prefix[STAT_PARAGRAPH_PREFIXES[3]])

    def table_number(text: str) -> float:
        return _number(text)

    top_net = [
        {
            "xmc": row.cells[0].text,
            "outflow": table_number(row.cells[1].text),
            "inflow": table_number(row.cells[2].text),
            "net": table_number(row.cells[3].text),
        }
        for row in source.tables[1].rows[1:]
    ]
    top_const = [
        {"xmc": row.cells[0].text, "n": int(table_number(row.cells[1].text)),
         "area": table_number(row.cells[2].text)}
        for row in source.tables[2].rows[1:]
    ]
    period_text = next(p.text for p in source.paragraphs if p.text.startswith("监测期："))
    period_dates = re.findall(r"(?<!年)\b\d{8}\b", period_text)
    return {
        "query_started_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "period": {"qsx": period_dates[0] if len(period_dates) > 1 else "—",
                   "hsx": period_dates[-1] if period_dates else "—"},
        "total_n": int(overall[0]), "total_area": overall[1],
        "orig_crop": {"n": int(crop[0]), "area": crop[1]},
        "cur_crop": {"n": int(crop[2]), "area": crop[3]},
        "crop_net_n": int(crop[4]), "crop_net_area": crop[5],
        "crop2const": {"n": int(flow[0]), "area": flow[1]},
        "restore2crop": {"n": int(flow[2]), "area": flow[3]},
        "cur_const": {"n": int(construction[0]), "area": construction[1]},
        "crop2const_pct": construction[4],
        "top_net_counties": top_net,
        "top_const_counties": top_const,
        "categories_display": [], "separate_types": [], "top_types": [],
    }


def hollow_report(source_path: Path, template_path: Path) -> dict[str, Any]:
    """把成品中的动态统计结果替换为槽位，返回可回填的标准答案清单。"""
    source_bytes = source_path.read_bytes()
    document = Document(source_path)
    manifest: dict[str, Any] = {
        "schema_version": 1,
        "source_file": source_path.name,
        "source_sha256": sha256(source_bytes).hexdigest(),
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "paragraph_answers": [], "table_answers": [],
    }
    for block_index, block in enumerate(_blocks(document)):
        if isinstance(block, Paragraph):
            if block.text.startswith("生成时间：") or block.text.startswith("　生成时间："):
                _set_paragraph_text(block, "　生成时间：[[SYSTEM_TIME]]")
                continue
            if not block.text.startswith(STAT_PARAGRAPH_PREFIXES):
                continue
            answers = NUMBER.findall(block.text)
            if not answers:
                continue
            manifest["paragraph_answers"].append({
                "block_index": block_index, "source_text": block.text,
                "answers": answers,
            })
            _set_paragraph_text(block, NUMBER.sub("***", block.text))

    for table_index, table in enumerate(document.tables):
        first_header = table.rows[0].cells[0].text if table.rows else ""
        first_dynamic = any(word in first_header for word in ("县", "市", "区", "行政区"))
        for row_index, row in enumerate(table.rows[1:], start=1):
            for column_index, cell in enumerate(row.cells):
                if column_index == 0 and not first_dynamic:
                    continue
                value = cell.text.strip()
                if not value:
                    continue
                manifest["table_answers"].append({
                    "table_index": table_index, "row": row_index,
                    "column": column_index, "answer": value,
                })
                _set_paragraph_text(cell.paragraphs[0], "***")

    template_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(template_path)
    return manifest


def refill_template(template_path: Path, manifest: dict[str, Any], output_path: Path) -> Path:
    """按稳定块序号和表格坐标回填，用于验证挖空过程是否丢失数据。"""
    document = Document(template_path)
    blocks = _blocks(document)
    for item in manifest["paragraph_answers"]:
        paragraph = blocks[item["block_index"]]
        if not isinstance(paragraph, Paragraph):
            raise ValueError("回填定位失效：目标不是段落。")
        answers = iter(item["answers"])
        restored = re.sub(r"\*{2,}", lambda _match: next(answers), paragraph.text)
        _set_paragraph_text(paragraph, restored)
    for item in manifest["table_answers"]:
        cell = document.tables[item["table_index"]].rows[item["row"]].cells[item["column"]]
        _set_paragraph_text(cell.paragraphs[0], item["answer"])
    for paragraph in document.paragraphs:
        if "[[SYSTEM_TIME]]" in paragraph.text:
            _set_paragraph_text(
                paragraph,
                paragraph.text.replace("[[SYSTEM_TIME]]", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M")),
            )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def clone_manifest(manifest: dict[str, Any]) -> dict[str, Any]:
    return deepcopy(manifest)
