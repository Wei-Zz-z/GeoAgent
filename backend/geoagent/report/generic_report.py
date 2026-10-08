"""用本轮已保存绑定生成普通模板报告，不调用专用模板生成器。"""

from datetime import datetime
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .result_binding import bind_charts
from .template_fill import refill_docx


def _align_report_elements(document: Any, parsed: dict[str, Any]) -> None:
    """仅整理报告中的表格和插图；不改正文内容、字体或统计绑定。"""
    table_blocks = [block for block in parsed["blocks"] if block["type"] == "table"]
    for table, block in zip(document.tables, table_blocks):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        if indent is not None:
            table._tbl.tblPr.remove(indent)
        header_rows = block.get("header_rows", 1)
        seen = set()
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    # 数值列靠右，名称靠左；表头居中。合并表头仍使用原有结构。
                    numeric = bool(re.fullmatch(r"[+-]?\d+(?:\.\d+)?%?", cell.text.strip()))
                    paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if row_index < header_rows else
                                           WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT)
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.left_indent = Pt(0)
                    paragraph.paragraph_format.right_indent = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    if row_index < header_rows:
                        paragraph.paragraph_format.keep_with_next = True
    titles = {block.get("table_title") for block in table_blocks}
    captions = {anchor["caption"] for anchor in parsed.get("chart_anchors", [])}
    for index, paragraph in enumerate(document.paragraphs):
        is_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        if paragraph.text in titles or paragraph.text in captions or is_drawing:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.right_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_with_next = paragraph.text in titles or is_drawing
            # 图形段落不继承正文固定行高，避免内嵌图被裁切。
            if is_drawing:
                paragraph.paragraph_format.line_spacing = 1
        if paragraph.text in titles and index + 1 < len(document.paragraphs):
            unit = document.paragraphs[index + 1]
            if unit.text.startswith("单位："):
                unit.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                unit.paragraph_format.first_line_indent = Pt(0)
                unit.paragraph_format.keep_with_next = True


def fill_generic_report(source: Path, parsed: dict[str, Any], output: Path,
                        scalars: dict[str, dict[str, Any]], tables: dict[str, dict[str, Any]],
                        started_at: datetime) -> dict[str, Any]:
    values = {slot["id"]: scalars.get(slot["id"], {}).get("value", "待核实")
              for slot in parsed["slots"]
              if slot["kind"] in {"placeholder", "dynamic_number", "dynamic_region"}}
    charts = bind_charts(parsed, tables)
    refill_docx(source, parsed, output, scalar_answers=values,
                table_answers={key: value["rows"] for key, value in tables.items()}, charts=charts,
                text_replacements={"[[SYSTEM_TIME]]": started_at.astimezone().strftime("%Y-%m-%d %H:%M")})
    pending = [slot["id"] for slot in parsed["slots"]
               if slot["id"] not in scalars and slot["id"] not in tables]
    pending.extend(anchor["id"] for anchor in parsed.get("chart_anchors", []) if anchor["id"] not in charts)
    reference_ids = {key for key, value in scalars.items() if value.get("status") == "仅供参考"}
    document = Document(output)
    _align_report_elements(document, parsed)
    # 保留原段落文字格式，参考段落额外使用底纹，不影响标准查询段落。
    for block in parsed["blocks"]:
        if any(slot["id"] in reference_ids and slot["block_id"] == block["id"] for slot in parsed["slots"]):
            # 插图会改变正文块序号，按原文固定部分查找替换后的段落。
            expected = block["text"]
            for slot in sorted((s for s in parsed["slots"] if s["block_id"] == block["id"]),
                               key=lambda s: s.get("start", 0), reverse=True):
                expected = expected[:slot["start"]] + str(values[slot["id"]]) + expected[slot["end"]:]
            for paragraph in document.paragraphs:
                if paragraph.text == expected:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "FFF2CC")
                    paragraph._p.get_or_add_pPr().append(shading)
                    paragraph.add_run("【未命中标准问题库，仅供参考】")
    if pending:
        pending_heading = document.add_heading("待核实事项", level=1)
        pending_heading.paragraph_format.page_break_before = True
        document.add_paragraph("以下内容尚未取得可靠结果或未完成绑定，不能作为已完成的统计结论：")
        for slot_id in pending:
            slot = next((item for item in parsed["slots"] if item["id"] == slot_id), None)
            if slot:
                block = next(item for item in parsed["blocks"] if item["id"] == slot["block_id"])
                label = block.get("table_title") or block.get("text") or slot_id
            else:
                label = next((item["caption"] for item in parsed.get("chart_anchors", []) if item["id"] == slot_id), slot_id)
            pending_paragraph = document.add_paragraph(label)
            pending_paragraph.paragraph_format.keep_together = True
    document.save(output)
    return {"filled_scalars": len(scalars), "filled_tables": len(tables),
            "charts": len(charts), "pending_items": pending,
            "chart_bindings": charts, "reference_slots": sorted(reference_ids)}
