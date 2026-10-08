"""标准问题库示例快报的模板制作与查询结果回填。"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..tools.labels import TBLX_LABELS
from .template_fill import refill_docx


REQUIRED_QUESTION_IDS = (
    "total_summary",
    "orig_farmland",
    "current_farmland",
    "by_type_summary",
    "by_county",
)


def is_question_library_template(parsed: dict[str, Any]) -> bool:
    """仅对已知示例模板开放无人值守回填，避免把相似问题误当作相同口径。"""
    structure = [
        ("p", block["text"]) if block["type"] == "paragraph" else (
            "t", block.get("table_title"),
            [header["label"] for header in block.get("column_headers", [])],
            len(block["rows"]),
        )
        for block in parsed["blocks"]
    ]
    digest = sha256(json.dumps(structure, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
    markers = [slot["text"] for slot in parsed["slots"] if slot["kind"] == "placeholder"]
    tables = [block for block in parsed["blocks"] if block["type"] == "table"]
    captions = [anchor["caption"] for anchor in parsed["chart_anchors"]]
    return (
        parsed.get("report_year") == 2026
        # 与 build_question_library_template 的正文和表头一致；改字句后需人工确认口径。
        and digest == "df65700cb93128cd2f5bf5f2165d8e7e2c614492209be2faf47b2bb6038e822a"
        and len(parsed["blocks"]) == 21
        and markers == [
            "{{total_n}}", "{{total_area_mu}}", "{{orig_farmland_n}}",
            "{{orig_farmland_area_mu}}", "{{current_farmland_n}}",
            "{{current_farmland_area_mu}}",
        ]
        and len(parsed["slots"]) == 8
        and [table.get("table_title") for table in tables] == [
            "表1 变化后图斑类型面积前10位", "表2 各县变化图斑面积前10位",
        ]
        and all(len(table["rows"]) == 11 for table in tables)
        and captions == ["图1 变化后主要图斑类型面积构成", "图2 各县变化图斑面积前10位"]
    )


def _set_font(run: Any, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def _add_paragraph(
    document: Document,
    text: str,
    *,
    heading: bool = False,
    center: bool = False,
) -> Any:
    paragraph = document.add_paragraph(style="Heading 1" if heading else None)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    if not heading:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    _set_font(
        paragraph.add_run(text),
        "黑体" if heading else "仿宋_GB2312",
        16,
        bold=heading,
    )
    return paragraph


def _shade_cell(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _style_table(table: Any, headers: list[str]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = headers[index]
        _shade_cell(cell, "D9EAF7")
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    _set_font(run, "黑体" if row_index == 0 else "宋体", 10.5,
                              bold=row_index == 0)


def build_question_library_template(path: str | Path) -> Path:
    """制作可被当前解析器识别、且仅使用标准问题库口径的示例快报。"""
    output = Path(path)
    document = Document()
    for style_name in ("Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        paragraph_properties = style._element.find(qn("w:pPr"))
        if paragraph_properties is not None:
            borders = paragraph_properties.find(qn("w:pBdr"))
            if borders is not None:
                paragraph_properties.remove(borders)
    section = document.sections[0]
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    _set_font(title.add_run("2026年第一期地类变化监测快报"), "方正小标宋简体", 22)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(14)
    _set_font(meta.add_run("查询时间：[[SYSTEM_TIME]]"), "楷体_GB2312", 14)

    _add_paragraph(document, "一、总体情况", heading=True)
    _add_paragraph(
        document,
        "2026年第一期全库共有{{total_n}}个变化图斑，总面积{{total_area_mu}}亩。"
        "本快报按照统一统计口径汇总变化前后地类信息，并对重点类型和行政区进行对比分析。",
    )

    _add_paragraph(document, "二、耕地变化情况", heading=True)
    _add_paragraph(
        document,
        "变化前耕地图斑{{orig_farmland_n}}个，总面积{{orig_farmland_area_mu}}亩。",
    )
    _add_paragraph(
        document,
        "变化后为耕地的图斑{{current_farmland_n}}个，总面积{{current_farmland_area_mu}}亩。",
    )

    _add_paragraph(document, "三、变化后图斑类型构成", heading=True)
    _add_paragraph(
        document,
        "按变化后图斑类型汇总数量和面积，面积排名前10位的统计结果见表1，"
        "主要类型的面积构成见图1。",
    )
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(caption.add_run("表1 变化后图斑类型面积前10位"), "黑体", 12, bold=True)
    unit = document.add_paragraph()
    unit.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(unit.add_run("单位：亩"), "宋体", 10.5)
    table = document.add_table(rows=11, cols=3)
    _style_table(table, ["图斑类型", "图斑数量", "总面积"])
    chart_caption = document.add_paragraph()
    chart_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(chart_caption.add_run("图1 变化后主要图斑类型面积构成"), "宋体", 10.5)

    _add_paragraph(document, "四、分县变化情况", heading=True)
    _add_paragraph(
        document,
        "按县级行政区汇总变化图斑数量和面积，面积排名前10位的统计结果见表2，"
        "各县变化面积对比见图2。",
    )
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(caption.add_run("表2 各县变化图斑面积前10位"), "黑体", 12, bold=True)
    unit = document.add_paragraph()
    unit.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(unit.add_run("单位：亩"), "宋体", 10.5)
    table = document.add_table(rows=11, cols=3)
    _style_table(table, ["县（市、区）", "图斑数量", "总面积"])
    chart_caption = document.add_paragraph()
    chart_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(chart_caption.add_run("图2 各县变化图斑面积前10位"), "宋体", 10.5)

    _add_paragraph(document, "五、统计说明", heading=True)
    _add_paragraph(
        document,
        "本快报面积由数据库平方米值换算为亩，换算关系为1平方米等于0.0015亩，"
        "结果保留两位小数。正文、表格和图表使用同一批查询结果，数字不使用千位分隔符。",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def _number(value: Any, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}"


def _entry_rows(executions: dict[str, dict[str, Any]], entry_id: str) -> list[dict[str, Any]]:
    result = executions.get(entry_id)
    if result is None:
        raise ValueError(f"标准问题尚未执行：{entry_id}")
    rows = result.get("rows")
    if not isinstance(rows, list) or not rows:
        raise ValueError(f"标准问题没有可回填结果：{entry_id}")
    return rows


def fill_question_library_report(
    template_path: str | Path,
    parsed: dict[str, Any],
    executions: dict[str, dict[str, Any]],
    output_path: str | Path,
    *,
    query_started_at: datetime,
) -> Path:
    """把五个标准问题的同批查询结果回填到示例模板。"""
    total = _entry_rows(executions, "total_summary")[0]
    original = _entry_rows(executions, "orig_farmland")[0]
    current = _entry_rows(executions, "current_farmland")[0]
    type_rows = _entry_rows(executions, "by_type_summary")[:10]
    county_rows = _entry_rows(executions, "by_county")[:10]
    marker_values = {
        "{{total_n}}": str(total["n"]),
        "{{total_area_mu}}": _number(total["area"] * 0.0015),
        "{{orig_farmland_n}}": str(original["n"]),
        "{{orig_farmland_area_mu}}": _number(original["area_m2"] * 0.0015),
        "{{current_farmland_n}}": str(current["n"]),
        "{{current_farmland_area_mu}}": _number(current["area_m2"] * 0.0015),
    }
    scalar_answers = {
        slot["id"]: marker_values[slot["text"]]
        for slot in parsed["slots"]
        if slot["kind"] == "placeholder" and slot["text"] in marker_values
    }
    table_answers: dict[str, list[list[Any]]] = {}
    for slot in parsed["slots"]:
        if slot["kind"] != "table_region":
            continue
        block = next(item for item in parsed["blocks"] if item["id"] == slot["block_id"])
        title = block.get("table_title") or block.get("caption", "")
        if "图斑类型" in title:
            table_answers[slot["id"]] = [
                [
                    TBLX_LABELS.get(str(row["TBLX"]), row["TBLX"]),
                    row["n"],
                    _number(row["area_m2"] * 0.0015),
                ]
                for row in type_rows
            ]
        elif "各县" in title:
            table_answers[slot["id"]] = [
                [row["region"], row["n"], _number(row["area"] * 0.0015)]
                for row in county_rows
            ]

    charts: dict[str, dict[str, Any]] = {}
    for anchor in parsed.get("chart_anchors", []):
        if "图斑类型" in anchor["caption"]:
            pie_rows = type_rows[:6]
            charts[anchor["id"]] = {
                "chart_type": "pie",
                "labels": [TBLX_LABELS.get(str(row["TBLX"]), row["TBLX"]) for row in pie_rows],
                "values": [row["area_m2"] * 0.0015 for row in pie_rows],
            }
        elif "各县" in anchor["caption"]:
            charts[anchor["id"]] = {
                "chart_type": "bar",
                "labels": [row["region"] for row in county_rows],
                "values": [row["area"] * 0.0015 for row in county_rows],
                "fill_color": "70AD47",
                "line_color": "548235",
            }
    output = refill_docx(
        template_path,
        parsed,
        output_path,
        scalar_answers=scalar_answers,
        table_answers=table_answers,
        charts=charts,
        text_replacements={
            "[[SYSTEM_TIME]]": query_started_at.astimezone().strftime("%Y-%m-%d %H:%M")
        },
    )
    document = Document(output)
    if len(document.tables) >= 2:
        _style_table(document.tables[0], ["图斑类型", "图斑数量", "总面积"])
        _style_table(document.tables[1], ["县（市、区）", "图斑数量", "总面积"])
    document.save(output)
    return output


def fill_review_report(
    template_path: str | Path,
    parsed: dict[str, Any],
    output_path: str | Path,
    *,
    query_started_at: datetime,
    confirmed_scalars: dict[str, str] | None = None,
    confirmed_table_cells: dict[str, str] | None = None,
    reference_only: bool = False,
) -> Path:
    """生成结构核验版；只回填已确认数值，其余明确标为待核验。"""
    scalar_answers = {
        slot["id"]: "待核验"
        for slot in parsed["slots"]
        if slot["kind"] in {"placeholder", "dynamic_region", "dynamic_number"}
    }
    scalar_answers.update(confirmed_scalars or {})
    output = refill_docx(
        template_path,
        parsed,
        output_path,
        scalar_answers=scalar_answers,
        table_cell_answers=confirmed_table_cells,
        test_notice=(
            "【核验版】自由问数的回填数值仅供参考，仍需业务核实；"
            if reference_only else
            "【核验版】仅已保存绑定的正文和简单表格数值按原位置回填；"
        ) + (
            "未绑定正文标记为“待核验”，未绑定表格保持原样。"
            "图表尚未绑定，不得作为正式业务成果使用。"
        ),
        text_replacements={
            "[[SYSTEM_TIME]]": query_started_at.astimezone().strftime("%Y-%m-%d %H:%M")
        },
        remove_chart_anchors=True,
    )
    confirmed_ids = set(confirmed_scalars or {})
    confirmed_cells = set(confirmed_table_cells or {})
    pending = []
    for slot in parsed["slots"]:
        if slot["kind"] in {"placeholder", "dynamic_region", "dynamic_number"}:
            if slot["id"] not in confirmed_ids:
                pending.append(f'{slot["id"]} · {slot.get("context", "正文数值")[:70]}')
        elif slot["kind"] == "table_region":
            target_cells = set(slot.get("target_cell_ids", []))
            if not target_cells or target_cells - confirmed_cells:
                pending.append(f'{slot["id"]} · 表格尚未完整核实')
    for anchor in parsed.get("chart_anchors", []):
        pending.append(f'{anchor["id"]} · 图表数据与位置尚未绑定')
    if pending:
        document = Document(output)
        document.add_heading("待核实事项", level=1)
        document.add_paragraph("以下内容尚未取得可靠答案或未完成位置绑定，不应作为正式业务结论。")
        for item in pending:
            document.add_paragraph(item, style="List Bullet")
        document.save(output)
    return output
