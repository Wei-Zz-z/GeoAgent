"""DOCX解析和上传接口回归，不依赖数据库或模型。"""
import base64
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
import pytest

from geoagent.report.template import decompose_atomic_items, decompose_questions, parse_docx
from geoagent.server.app import create_app


def fixture_docx(text: str = '耕地流入***亩') -> bytes:
    document = Document()
    paragraph = document.add_paragraph(style='Heading 1')
    paragraph.add_run(text[:2]).bold = True
    paragraph.add_run(text[2:])
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = '地区'
    table.cell(0, 1).text = '流入'
    table.cell(0, 2).text = '流出'
    table.cell(1, 0).text = '杭州市'
    table.cell(2, 0).text = '宁波市'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def merged_header_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=4, cols=4)
    table.cell(0, 0).merge(table.cell(1, 0)).text = '行政区划'
    table.cell(0, 1).merge(table.cell(0, 2)).text = '总计'
    table.cell(0, 3).text = '其中'
    table.cell(1, 1).text = '面积'
    table.cell(1, 2).text = '占比'
    table.cell(1, 3).text = '耕地'
    table.cell(2, 0).text = '浙江省'
    table.cell(3, 0).text = '杭州市'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def titled_table_docx() -> bytes:
    document = Document()
    document.add_paragraph('表1 各市耕地净变化情况表')
    document.add_paragraph('单位：亩')
    table = document.add_table(rows=3, cols=4)
    for index, text in enumerate(('行政区划', '流出面积', '流入面积', '净变化面积')):
        table.cell(0, index).text = text
    table.cell(1, 0).text = '浙江省'
    table.cell(2, 0).text = '杭州市'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def figure_docx() -> bytes:
    document = Document()
    document.add_paragraph('一、耕地变化情况')
    image = BytesIO(base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='
    ))
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(image)
    document.add_paragraph('图1 耕地流入和流出情况')
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_parse_and_question_provenance() -> None:
    parsed = parse_docx(fixture_docx(), '模板.docx')
    assert len(parsed['slots']) == 2
    assert parsed['slots'][0]['text'] == '***'
    assert parsed['slots'][1]['kind'] == 'table_region'
    assert len(parsed['slots'][1]['target_cell_ids']) == 4
    assert parsed['blocks'][0]['style']['heading_level'] == 1
    assert parsed['blocks'][0]['runs'][0]['bold'] is True
    questions = decompose_questions(parsed)
    assert len(questions) == 2
    assert '耕地流入' in questions[0]['question']
    assert questions[0]['slot_ids'] == ['b0:s0']
    assert all(q['status'] == '待用户确认' for q in questions)
    assert questions[0]['source_text'] == '耕地流入***亩'
    assert questions[0]['source_locator']['part'] == 'word/document.xml'


def test_questions_change_with_document() -> None:
    a = decompose_questions(parse_docx(fixture_docx('林地面积*亩'), '甲.docx'))
    b = decompose_questions(parse_docx(fixture_docx('耕地面积*亩'), '乙.docx'))
    assert a[0]['question'] != b[0]['question']


def test_sections_atomic_items_and_match_stub() -> None:
    parsed = parse_docx(fixture_docx('全省耕地流入***万亩，流出***万亩'), '模板.docx')
    assert parsed['schema_version'] == 4
    assert parsed['sections'][0]['heading_block'] == 'b0'
    assert parsed['blocks'][0]['section_id'] == parsed['sections'][0]['id']
    items = decompose_atomic_items(parsed)
    scalar_items = [item for item in items if item['kind'] == 'scalar']
    assert len(scalar_items) == 2
    assert scalar_items[0]['metric'].endswith('面积')
    assert scalar_items[0]['binding']['slot_id'] == 'b0:s0'
    assert scalar_items[0]['question_library_match']['status'] == 'unmatched'
    questions = decompose_questions(parsed)
    assert questions[0]['atomic_item_ids'] == [item['id'] for item in scalar_items]
    assert questions[0]['question_library_match']['candidates'] == []


def test_chart_caption_links_existing_image() -> None:
    parsed = parse_docx(figure_docx(), '图表模板.docx')
    assert len(parsed['chart_anchors']) == 1
    anchor = parsed['chart_anchors'][0]
    assert anchor['caption'] == '图1 耕地流入和流出情况'
    assert anchor['placement'] == 'replace_existing_image'
    assert anchor['target_image_id'] == parsed['images'][0]['id']
    assert anchor['relationship_id']


def test_parallel_placeholders_become_distinct_metrics() -> None:
    text = ('全省耕地流入***万亩，其中耕地流入中，来源于建设用地***万亩'
            '以及坑塘水面***万亩。')
    items = [item for item in decompose_atomic_items(parse_docx(fixture_docx(text), '模板.docx'))
             if item['kind'] == 'scalar']
    assert [item['metric'] for item in items] == [
        '耕地流入面积', '耕地流入来源于建设用地面积', '耕地流入来源于坑塘水面面积',
    ]
    assert len({item['binding']['slot_id'] for item in items}) == 3


def test_ranked_region_and_area_are_related() -> None:
    text = '新增违法面积大于1000亩的县（市、区）有**个，其中**县为***亩。'
    items = [item for item in decompose_atomic_items(parse_docx(fixture_docx(text), '模板.docx'))
             if item['kind'] == 'scalar']
    assert items[0]['expected_type'] == 'integer'
    assert items[1]['expected_type'] == 'region_name'
    assert items[2]['related_to'] == items[1]['id']
    assert items[2]['metric'] == items[1]['metric'] + '面积'


def test_question_removes_presentation_reference_and_historical_region() -> None:
    text = '净减少最多的为温州市，净减少***亩，具体详见图1；排名如表2所示；分市统计情况详见表3。'
    question = decompose_questions(parse_docx(fixture_docx(text), '模板.docx'))[0]
    assert '图1' not in question['question']
    assert '表2' not in question['question']
    assert '表3' not in question['question']
    assert '温州市' not in question['question']
    assert '[待查询地区]' in question['question']
    assert question['historical_examples'] == ['温州市']
    assert '不得作为筛选条件' in question['requirements']


def test_historical_rank_region_becomes_refill_anchor() -> None:
    text = ('全省耕地净减少的地市有*个，其中净减少最多的为温州市，净减少***万亩；'
            '全省耕地净减少的县（市、区）有***个，其中净减少最多的为义乌市，'
            '净减少***万亩。')
    parsed = parse_docx(fixture_docx(text), '模板.docx')
    items = [item for item in decompose_atomic_items(parsed) if item['kind'] == 'scalar']
    metrics = [item['metric'] for item in items]
    assert '耕地净减少设区市数量' in metrics
    assert '耕地净减少县级行政区数量' in metrics
    assert '耕地净减少最多设区市' in metrics
    assert '耕地净减少最多县级行政区' in metrics
    assert any(slot['kind'] == 'dynamic_region' and slot['text'] == '温州市'
               for slot in parsed['slots'])


def test_merged_header_becomes_column_paths() -> None:
    parsed = parse_docx(merged_header_docx(), '模板.docx')
    table = next(block for block in parsed['blocks'] if block['type'] == 'table')
    assert table['header_rows'] == 2
    assert [item['label'] for item in table['column_headers']] == [
        '行政区划', '总计 / 面积', '总计 / 占比', '地类构成 / 耕地',
    ]
    question = decompose_questions(parsed)[0]
    assert '总计 / 面积' in question['source_text']
    assert question['source_text'].count('行政区划') == 1


def test_table_uses_nearby_title_and_unit_as_atomic_metric() -> None:
    parsed = parse_docx(titled_table_docx(), '模板.docx')
    table = next(block for block in parsed['blocks'] if block['type'] == 'table')
    assert table['table_title'] == '表1 各市耕地净变化情况表'
    assert table['unit'] == '亩'
    item = next(item for item in decompose_atomic_items(parsed) if item['kind'] == 'table')
    assert item['metric'] == '各市耕地净变化情况表'
    assert item['unit'] == '亩'


def test_user_question_uses_2026_business_period_but_keeps_source_text() -> None:
    parsed = parse_docx(fixture_docx('2024年上半年耕地流入***亩'), '模板.docx')
    question = decompose_questions(parsed)[0]['question']
    assert parsed['blocks'][0]['text'] == '2024年上半年耕地流入***亩'
    assert '2026年第一期' in question
    assert '2024年' not in question
    assert parsed['report_year'] == 2026
    assert parsed['parsed_at']


@pytest.mark.parametrize('name,data', [('模板.doc', b'abc'), ('模板.docx', b'abc')])
def test_invalid_upload(name: str, data: bytes) -> None:
    with pytest.raises(ValueError):
        parse_docx(data, name)


def test_entity_rejected() -> None:
    data = BytesIO()
    with ZipFile(data, 'w') as package:
        package.writestr('word/document.xml', '<!DOCTYPE a [<!ENTITY x SYSTEM "file:///secret">]><a/>')
    with pytest.raises(ValueError, match='XML'):
        parse_docx(data.getvalue(), '模板.docx')


def test_upload_endpoint(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv('GEOAGENT_DATA_DIR', str(tmp_path))
    with TestClient(create_app()) as client:
        response = client.post('/api/report-templates/parse?filename=template.docx', content=fixture_docx())
        assert response.status_code == 200
        assert response.json()['stage'] == '请确认问数清单'
        assert response.json()['atomic_items']
        assert client.post('/api/report-templates/parse?filename=bad.docx', content=b'bad').status_code == 400
