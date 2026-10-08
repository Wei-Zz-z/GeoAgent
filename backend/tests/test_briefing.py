"""快报生成测试：TBLX 三大类映射与 docx 构建（不依赖真实库）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from geoagent.report.briefing import build_briefing_docx
from geoagent.tools.categories import (
    category_codes,
    load_entries,
    tblx_category,
)
from geoagent.tools.labels import TBLX_LABELS


def test_tblx_category_covers_all_label_codes():
    assert set(load_entries()) == set(TBLX_LABELS)


def test_tblx_category_key_types():
    assert tblx_category("01") == "agricultural"
    assert tblx_category("1") == "agricultural"
    assert tblx_category("20") == "construction"
    assert tblx_category("DT") == "separate"
    assert tblx_category("HL") == "unused"
    assert tblx_category("NOPE") is None


def test_category_codes_include_zero_padded_variants():
    construction = set(category_codes("construction"))
    assert "01" not in construction
    assert {"20", "YH", "DL1"} <= construction
    assert not {"DT", "TD", "WL", "TP", "QT"} & construction
    agricultural = set(category_codes("agricultural"))
    assert {"01", "1", "03", "3", "KT", "SK"} <= agricultural


def _fake_stats() -> dict:
    county = lambda xmc, outflow, inflow: {  # noqa: E731
        "xmc": xmc,
        "outflow": outflow,
        "inflow": inflow,
        "net": inflow - outflow,
    }
    const = lambda xmc, n, area: {"xmc": xmc, "n": n, "area": area}  # noqa: E731
    return {
        "query_started_at": "2026-09-15T19:30:45+08:00",
        "period": {"qsx": "20251006", "hsx": "20260501"},
        "total_n": 164798,
        "total_area": 20586122.73,
        "orig_crop": {"n": 46880, "area": 5847014.45},
        "cur_crop": {"n": 13901, "area": 1736466.47},
        "crop2const": {"n": 1234, "area": 150000.0},
        "restore2crop": {"n": 500, "area": 60000.0},
        "cur_const": {"n": 23591, "area": 2954687.33},
        "crop_net_n": 13901 - 46880,
        "crop_net_area": 1736466.47 - 5847014.45,
        "crop2const_pct": 150000.0 / 2954687.33 * 100,
        "top_net_counties": [
            county("义乌市", 100000, 20000),
            county("温州市", 90000, 30000),
        ],
        "top_const_counties": [
            const("义乌市", 100, 200000.0),
            const("萧山区", 90, 180000.0),
        ],
        "categories_display": [
            ("农用地", [("01", "耕地", False), ("ND", "农村道路", True)]),
            ("建设用地", [("20", "建/构筑物", False)]),
        ],
    }


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def test_build_briefing_docx_writes_report(tmp_path):
    out = build_briefing_docx(_fake_stats(), tmp_path / "快报.docx")
    assert out.exists()
    text = _docx_text(out)
    assert "主要地类变化监测快报" in text
    assert "164798" in text
    assert "30879.18" in text
    assert "图斑数（个）" not in text
    assert "-32,979" not in text
    assert "面积（平方米）" not in text
    assert "义乌市" in text
    assert "疑似违法占地" in text
    assert "附录　图斑类型→三大类" in text
    assert "农村道路(ND)（待确认）" in text
    assert "查询时间：2026-09-15 19:30" in text


def test_briefing_formatting_matches_official_template(tmp_path):
    """排版回归：字体与段落格式对齐《快报模板（删除统计数值）.doc》。"""
    out = build_briefing_docx(_fake_stats(), tmp_path / "快报格式.docx")
    doc = Document(str(out))

    sec = doc.sections[0]
    assert round(sec.page_width.mm, 1) == 210.0
    assert round(sec.page_height.mm, 1) == 297.0
    assert round(sec.left_margin.cm, 2) == 3.0
    assert round(sec.right_margin.cm, 2) == 3.0
    assert round(sec.top_margin.cm, 2) == 2.54
    assert round(sec.bottom_margin.cm, 2) == 2.54

    def run_fonts(p):
        r = p.runs[0]
        rPr = r._element.rPr
        rFonts = rPr.rFonts if rPr is not None else None
        return {
            "size": r.font.size.pt if r.font.size else None,
            "east_asia": rFonts.get(qn("w:eastAsia")) if rFonts is not None else None,
            "ascii": rFonts.get(qn("w:ascii")) if rFonts is not None else None,
            "bold": r.font.bold,
        }

    title = next(p for p in doc.paragraphs if p.text == "主要地类变化监测快报")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert run_fonts(title) == {
        "size": 22,
        "east_asia": "方正小标宋简体",
        "ascii": "方正小标宋简体",
        "bold": None,
    }

    meta = next(p for p in doc.paragraphs if p.text.startswith("监测期："))
    assert meta.alignment == WD_ALIGN_PARAGRAPH.CENTER
    fonts = run_fonts(meta)
    assert fonts["size"] == 20
    assert fonts["east_asia"] == "楷体_GB2312"
    assert fonts["ascii"] == "Times New Roman"

    query_time = next(p for p in doc.paragraphs if p.text.startswith("查询时间："))
    assert query_time.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert run_fonts(query_time)["east_asia"] == "楷体_GB2312"

    heading = next(p for p in doc.paragraphs if p.text.startswith("一、总体情况"))
    assert run_fonts(heading)["east_asia"] == "黑体"
    assert run_fonts(heading)["size"] == 16
    assert run_fonts(heading)["bold"] is not True
    pPr = heading._p.pPr
    ind = pPr.find(qn("w:ind"))
    assert ind is not None and ind.get(qn("w:firstLineChars")) == "200"

    body = next(p for p in doc.paragraphs if p.text.startswith("本期通过遥感影像"))
    assert run_fonts(body)["east_asia"] == "仿宋_GB2312"
    assert run_fonts(body)["size"] == 16
    spacing = body._p.pPr.find(qn("w:spacing"))
    assert spacing is not None
    assert spacing.get(qn("w:line")) == "360"
    assert spacing.get(qn("w:lineRule")) == "auto"

    table = doc.tables[0]
    header = table.rows[0].cells[0]
    shd = header._tc.tcPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "D9D9D9"
    assert run_fonts(header.paragraphs[0])["bold"] is True
    assert run_fonts(header.paragraphs[0])["east_asia"] == "仿宋_GB2312"


def test_word_contains_four_chart_parts(tmp_path):
    from zipfile import ZipFile
    from lxml import etree
    stats = _fake_stats()
    stats['top_types'] = [{'code': '01', 'area': 1000}, {'code': 'DT', 'area': 2000}]
    stats['separate_types'] = [{'code': 'DT', 'area': 2000}]
    path = build_briefing_docx(stats, tmp_path / 'charts.docx')
    text = _docx_text(path)
    assert '表1　耕地变化总体情况' in text
    assert '图1　变化后图斑类型面积前10位（单位：亩）' in text
    assert '图2　耕地净减少县域面积排名（单位：亩）' in text
    assert '图3　变化后建设用地县域面积排名（单位：亩）' in text
    assert '图4　动土等其他类型单列面积（单位：亩）' in text
    with ZipFile(path) as package:
        charts = sorted(n for n in package.namelist() if n.startswith('word/charts/'))
        assert len(charts) == 4
        expected_colors = ['4472C4', 'ED7D31', '70AD47', '8064A2']
        for name, expected_color in zip(charts, expected_colors):
            xml = etree.fromstring(package.read(name))
            assert xml.xpath('count(//*[local-name()="barChart"])') == 1
            assert xml.xpath('string(//*[local-name()="invertIfNegative"]/@val)') == '0'
            assert xml.xpath('string(//*[local-name()="ser"]/*[local-name()="spPr"]//*[local-name()="srgbClr"]/@val)') == expected_color
