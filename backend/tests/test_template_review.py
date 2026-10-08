"""模板完整性与模型复核接口测试，不访问真实模型。"""

from io import BytesIO
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient

from geoagent.report.review import audit_template_mapping
from geoagent.report.template import decompose_atomic_items, decompose_questions, parse_docx
from geoagent.server.app import create_app


def _template() -> bytes:
    document = Document()
    document.add_paragraph("全省耕地流入***亩")
    document.add_paragraph("请分析本期耕地流出的主要去向。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_audit_identifies_missing_binding() -> None:
    parsed = parse_docx(_template(), "测试.docx")
    atomic = decompose_atomic_items(parsed)
    questions = decompose_questions(parsed)
    audit = audit_template_mapping(parsed, atomic, questions)
    assert audit["covered_slot_count"] == audit["slot_count"] == 1
    assert audit["covered_atomic_item_count"] == 1
    questions[0]["slot_ids"] = []
    audit = audit_template_mapping(parsed, atomic, questions)
    assert any(issue["message"] == "待填位置没有对应问题" for issue in audit["issues"])


def test_model_review_keeps_suggestions_separate(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path))
    app = create_app()

    async def fake_chat(**kwargs):
        prompt = kwargs["messages"][0]["content"]
        assert "全省耕地流入" in prompt
        assert "耕地流出的主要去向" in prompt
        assert '"has_question": false' in prompt
        return SimpleNamespace(content=(
            '{"summary":"有一项口径需确认","items":[{"question_id":"Q1",'
            '"status":"需复核","reason":"需要确认流入定义","missing_metrics":[]}],'
            '"missing_items":[]}'
        ))

    app.state.llm.chat = fake_chat
    with TestClient(app) as client:
        response = client.post("/api/report-templates/parse?filename=测试.docx", content=_template())
        assert response.status_code == 200
        payload = response.json()
        review = client.post(
            f'/api/report-templates/{payload["workflow_id"]}/review',
            json={"questions": {"Q1": payload["questions"][0]["question"]}},
        )
        assert review.status_code == 200
        assert review.json()["items"][0]["status"] == "需复核"
        assert payload["questions"][0]["status"] == "待用户确认"
