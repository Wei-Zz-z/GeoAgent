"""简单固定地区行表格的人工绑定与原位回填。"""

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from geoagent.report.table_binding import fixed_row_table_options
from geoagent.report.template import parse_docx
from geoagent.server.app import create_app


def _template(title: str = "各市变化情况表") -> bytes:
    document = Document()
    document.add_paragraph(title)
    table = document.add_table(rows=3, cols=3)
    for index, value in enumerate(("行政区划", "流入面积", "流出面积")):
        table.cell(0, index).text = value
    table.cell(1, 0).text = "合计"
    table.cell(2, 0).text = "杭州市"
    run = table.cell(1, 1).paragraphs[0].add_run("")
    run.bold = True
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_ranked_table_not_treated_as_fixed_rows() -> None:
    content = _template("净减少最多的前10个地区")
    parsed = parse_docx(content, "排名.docx")
    assert fixed_row_table_options(parsed) == []


def test_fixed_table_binding_preserves_row_names_and_style(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))

    def fake_pdf(_source: Path, output: Path) -> Path:
        output.write_bytes(b"%PDF-1.4\n%%EOF")
        return output

    monkeypatch.setattr("geoagent.server.template_routes.convert_docx_to_pdf", fake_pdf)
    app = create_app()
    with TestClient(app) as client:
        parsed = client.post("/api/report-templates/parse?filename=简单表.docx", content=_template())
        assert parsed.status_code == 200
        payload = parsed.json()
        option = payload["fixed_tables"][0]
        assert [row["key"] for row in option["rows"]] == ["合计", "杭州市"]
        first_cell = option["rows"][0]["cells"][0]["id"]
        second_cell = option["rows"][1]["cells"][0]["id"]
        url = f'/api/report-templates/{payload["workflow_id"]}/bindings/fixed-table'
        body = {
            "slot_id": option["slot_id"], "source_question_id": payload["questions"][0]["id"],
            "source_note": "Q1问数结果，单位亩",
            "cells": [{"cell_id": first_cell, "value": "12.50"},
                      {"cell_id": second_cell, "value": "3.25"}],
        }
        assert client.post(url, json={**body, "cells": [{"cell_id": "b99:r1:c1", "value": "1"}]}).status_code == 400
        assert client.post(url, json={**body, "cells": [{"cell_id": first_cell, "value": "1,200"}]}).status_code == 400
        saved = client.post(url, json=body)
        assert saved.status_code == 200
        generated = client.post(
            f'/api/report-templates/{payload["workflow_id"]}/generate',
            json={"binding_mode": True, "confirmed_table_slot_ids": [option["slot_id"]]},
        )
        assert generated.status_code == 200
        document = Document(tmp_path / "reports" / generated.json()["word"]["filename"])
        table = document.tables[0]
        assert [table.cell(row, 0).text for row in (1, 2)] == ["合计", "杭州市"]
        assert table.cell(1, 1).text == "12.50"
        assert table.cell(2, 1).text == "3.25"
        assert table.cell(1, 2).text == ""
        assert table.cell(1, 1).paragraphs[0].runs[0].bold is True
