from pathlib import Path

from docx import Document

from geoagent.report.reverse_template import hollow_report, refill_template
from geoagent.report.template import decompose_questions, parse_docx


def _sample(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("生成时间：2026-09-06 17:06")
    doc.add_paragraph("一、总体情况")
    doc.add_paragraph("本期通过遥感影像与国土变更调查数据对比，共提取主要地类新发生变化图斑 100 个，总面积 200.50 平方米。")
    doc.add_paragraph("表2　县域排名")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "县（市、区）"
    table.rows[0].cells[1].text = "面积（平方米）"
    table.rows[1].cells[0].text = "义乌市"
    table.rows[1].cells[1].text = "200.50"
    doc.save(path)
    return path


def test_hollow_parse_and_refill_roundtrip(tmp_path: Path) -> None:
    source = _sample(tmp_path / "source.docx")
    template = tmp_path / "template.docx"
    restored = tmp_path / "restored.docx"
    manifest = hollow_report(source, template)
    parsed = parse_docx(template.read_bytes(), template.name)
    questions = decompose_questions(parsed)
    assert len(questions) == 2
    assert "***" in Document(template).paragraphs[2].text
    table_slot = next(slot for slot in parsed["slots"] if slot["kind"] == "table_region")
    assert "b4:r1:c0" in table_slot["target_cell_ids"]
    refill_template(template, manifest, restored)
    restored_doc = Document(restored)
    assert "100 个" in restored_doc.paragraphs[2].text
    assert restored_doc.tables[0].rows[1].cells[0].text == "义乌市"
