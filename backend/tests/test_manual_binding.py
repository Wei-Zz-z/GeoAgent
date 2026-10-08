"""非标准模板的正文数值人工绑定闭环。"""

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from geoagent.server.app import create_app


def _template() -> bytes:
    document = Document()
    document.add_paragraph("本期耕地流入***亩，流出***亩，净变化***亩。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_confirm_scalar_bindings_and_generate_review_docx(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))

    def fake_pdf(_source: Path, output: Path) -> Path:
        output.write_bytes(b"%PDF-1.4\n%%EOF")
        return output

    monkeypatch.setattr("geoagent.server.template_routes.convert_docx_to_pdf", fake_pdf)
    app = create_app()
    with TestClient(app) as client:
        parsed = client.post("/api/report-templates/parse?filename=非标准.docx", content=_template())
        assert parsed.status_code == 200
        payload = parsed.json()
        workflow_id = payload["workflow_id"]
        question = payload["questions"][0]
        slots = [slot for slot in payload["template"]["slots"] if slot["kind"] == "placeholder"]
        assert len(slots) == 3

        def bind(slot_id: str, value: str, note: str = "Q1智能体回答"):
            return client.post(
                f"/api/report-templates/{workflow_id}/bindings/scalar",
                json={"slot_id": slot_id, "value": value,
                      "source_question_id": question["id"], "source_note": note},
            )

        assert bind(slots[0]["id"], "1,234").status_code == 400
        assert bind(slots[0]["id"], "123.45", "").status_code == 400
        assert bind("b99:s0", "123.45").status_code == 400
        assert bind(slots[0]["id"], "123.45").status_code == 200
        assert bind(slots[1]["id"], "67.8").status_code == 200

        generated = client.post(
            f"/api/report-templates/{workflow_id}/generate",
            json={"binding_mode": True, "confirmed_slot_ids": [slots[0]["id"], slots[1]["id"]]},
        )
        assert generated.status_code == 200
        result = generated.json()
        assert result["review_only"] is True
        document = Document(tmp_path / "reports" / result["word"]["filename"])
        assert "【核验版】" in document.paragraphs[0].text
        assert document.paragraphs[1].text == "本期耕地流入123.45亩，流出67.8亩，净变化待核验亩。"
        assert client.get(result["word"]["url"]).status_code == 200
