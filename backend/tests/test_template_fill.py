from io import BytesIO

from docx import Document
from docx.shared import Pt
import pytest

from geoagent.report.template import parse_docx
from geoagent.report.template_fill import refill_docx


def test_refill_keeps_paragraph_run_style_and_fills_table(tmp_path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('总数').bold = True
    paragraph.add_run('{{total_n}}')
    paragraph.add_run('个')
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = '地区'
    table.cell(0, 1).text = '面积'
    source = BytesIO()
    document.save(source)
    template = tmp_path / 'template.docx'
    template.write_bytes(source.getvalue())
    parsed = parse_docx(source.getvalue(), template.name)
    scalar = next(slot for slot in parsed['slots'] if slot['kind'] == 'placeholder')
    table_slot = next(slot for slot in parsed['slots'] if slot['kind'] == 'table_region')
    output = tmp_path / 'filled.docx'
    refill_docx(
        template,
        parsed,
        output,
        scalar_answers={scalar['id']: 12},
        table_answers={table_slot['id']: [['杭州市', '1.20'], ['宁波市', '0.80']]},
        text_replacements={'总数': '本期总数'},
    )
    result = Document(output)
    assert result.paragraphs[0].text == '本期总数12个'
    assert result.paragraphs[0].runs[0].bold is True
    assert result.tables[0].cell(1, 0).text == '杭州市'


def test_table_fill_keeps_cell_style_when_adding_rows(tmp_path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '地区'
    table.cell(0, 1).text = '面积'
    table.cell(1, 0).text = '待填'
    run = table.cell(1, 1).paragraphs[0].add_run('***')
    run.font.size = Pt(12)
    run.bold = True
    stream = BytesIO()
    document.save(stream)
    source = tmp_path / 'source.docx'
    source.write_bytes(stream.getvalue())
    parsed = parse_docx(stream.getvalue(), source.name)
    slot = next(slot for slot in parsed['slots'] if slot['kind'] == 'table_region')
    output = tmp_path / 'result.docx'
    refill_docx(source, parsed, output, scalar_answers={}, table_answers={slot['id']: [
        ['甲县', '1.25'], ['乙县', '2.50'],
    ]})
    result = Document(output).tables[0]
    assert result.cell(2, 0).text == '乙县'
    assert result.cell(1, 1).paragraphs[0].runs[0].bold is True
    assert result.cell(2, 1).paragraphs[0].runs[0].font.size.pt == 12


def test_refill_rejects_changed_template(tmp_path) -> None:
    first = tmp_path / 'source.docx'
    document = Document()
    document.add_paragraph('面积***亩')
    stream = BytesIO()
    document.save(stream)
    first.write_bytes(stream.getvalue())
    parsed = parse_docx(stream.getvalue(), first.name)
    document.add_paragraph('后加的一行')
    document.save(first)
    with pytest.raises(ValueError, match='重新上传'):
        refill_docx(first, parsed, tmp_path / 'out.docx', scalar_answers={})
