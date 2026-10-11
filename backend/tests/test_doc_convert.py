"""老式 .doc 模板转换与 Word 样式兜底测试（不依赖真实 Word / 网络）。"""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from geoagent.report import docconvert
from geoagent.report.docx_styles import add_bullet_safe, add_heading_safe
from geoagent.server.app import create_app

OLE2_HEADER = bytes.fromhex("d0cf11e0a1b11ae1") + b"\x00" * 32


def _docx_bytes(text: str = "本期变化图斑 *** 个，面积 *** 亩。") -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def test_is_legacy_doc_detects_ole2():
    assert docconvert.is_legacy_doc(OLE2_HEADER) is True
    assert docconvert.is_legacy_doc(_docx_bytes()) is False
    assert docconvert.looks_like_docx(_docx_bytes()) is True
    assert docconvert.looks_like_docx(OLE2_HEADER) is False


def test_ensure_docx_passes_through_docx():
    payload = _docx_bytes()
    data, name = docconvert.ensure_docx(payload, "快报模板.docx")
    assert data == payload
    assert name == "快报模板.docx"
    # 扩展名不规范时统一成 .docx
    _, renamed = docconvert.ensure_docx(payload, "快报模板")
    assert renamed == "快报模板.docx"


def test_ensure_docx_converts_legacy_doc(monkeypatch):
    converted = _docx_bytes("转换后的模板 *** 亩")

    def fake_convert(src: Path, out_dir: Path) -> Path:
        assert src.read_bytes()[:8] == docconvert.OLE2_MAGIC
        target = out_dir / "source.docx"
        target.write_bytes(converted)
        return target

    monkeypatch.setattr(docconvert, "_convert_doc_to_docx_file", fake_convert)
    data, name = docconvert.ensure_docx(OLE2_HEADER, "快报模板（删除统计数值）.doc")
    assert data == converted
    assert name == "快报模板（删除统计数值）.docx"


def test_ensure_docx_rejects_other_formats():
    with pytest.raises(docconvert.DocConversionError):
        docconvert.ensure_docx(b"%PDF-1.4 fake", "模板.pdf")


def test_convert_reports_clear_error_without_converter(monkeypatch):
    """两种转换器都不可用时，报错要说清楚怎么自救。"""

    def no_word(src: Path, dst: Path) -> Path:
        raise docconvert.DocConversionError("Word 不可用")

    monkeypatch.delenv("GEOAGENT_DOC_CONVERTER", raising=False)
    monkeypatch.setattr(docconvert.shutil, "which", lambda _name: None)
    monkeypatch.setattr(docconvert, "_convert_with_word", no_word)
    with pytest.raises(docconvert.DocConversionError) as excinfo:
        docconvert.ensure_docx(OLE2_HEADER, "模板.doc")
    assert "另存为 .docx" in str(excinfo.value)


def test_add_bullet_safe_falls_back_when_style_missing():
    class FakeDocument:
        def __init__(self) -> None:
            self.texts: list[str] = []

        def add_paragraph(self, text=None, style=None):
            if style is not None:
                raise KeyError(f"no style with name '{style}'")
            self.texts.append(text)
            return SimpleNamespace(runs=[], text=text)

        def add_heading(self, text, level=1):
            raise KeyError("no style with name 'Heading 1'")

    document = FakeDocument()
    add_bullet_safe(document, "第一条")
    add_heading_safe(document, "待核实事项")

    assert document.texts == ["· 第一条", "待核实事项"]


def test_upload_accepts_legacy_doc(monkeypatch, tmp_path):
    """上传 .doc 时自动转成 .docx 再解析（转换器用桩替换，不启动 Word）。"""
    converted = _docx_bytes()

    def fake_convert(src: Path, out_dir: Path) -> Path:
        target = out_dir / "source.docx"
        target.write_bytes(converted)
        return target

    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(docconvert, "_convert_doc_to_docx_file", fake_convert)
    app = create_app()
    with TestClient(app) as client:
        response = client.post(
            "/api/report-templates/parse",
            params={"filename": "快报模板（删除统计数值）.doc"},
            content=OLE2_HEADER,
        )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["questions"]
    # 转换后的 docx 才是工作流的解析/回填对象
    stored = app.state.template_workflows[payload["workflow_id"]]["filename"]
    assert stored == "快报模板（删除统计数值）.docx"


def test_upload_rejects_unparseable_doc(monkeypatch, tmp_path):
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))

    def failing_convert(src: Path, out_dir: Path) -> Path:
        raise docconvert.DocConversionError("转换器不可用")

    monkeypatch.setattr(docconvert, "_convert_doc_to_docx_file", failing_convert)
    app = create_app()
    with TestClient(app) as client:
        response = client.post(
            "/api/report-templates/parse",
            params={"filename": "坏模板.doc"},
            content=OLE2_HEADER,
        )

    assert response.status_code == 400
    assert "转换" in response.json()["detail"]
