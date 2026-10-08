"""甲方固定快报试运行：固定查询、自动回填及缺失数据提示。"""

from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
import pytest

from geoagent.report.client_trial import fill_client_trial_report, query_client_trial
from geoagent.report.template import parse_docx
from geoagent.server.app import create_app


class _FakePg:
    async def run_sql(self, sql: str, **_: object) -> dict:
        if "AS from_orchard" in sql:
            return {"rows": [{
                "inflow": 40000, "outflow": 60000, "to_construction": 10000,
                "from_orchard": 12000, "from_forest": 8000,
                "from_construction": 2000, "from_pond": 1000,
                "construction": 50000,
            }]}
        if 'LEFT("XZQDM"::text,4)' in sql:
            return {"rows": [
                {"code": "3301", "inflow": 10000, "outflow": 30000},
                {"code": "3302", "inflow": 30000, "outflow": 30000},
            ]}
        return {"rows": [
            {"code": "330101", "name": "甲区", "inflow": 10000, "outflow": 30000},
            {"code": "330201", "name": "乙区", "inflow": 30000, "outflow": 30000},
        ]}

    async def close(self) -> None:
        return None


def _template() -> bytes:
    doc = Document()
    doc.add_paragraph("2024年上半年主要地类变化监测")
    doc.add_paragraph("（2024年6月24日）")
    doc.add_paragraph("一、全省耕地数量稳定且略有增加")
    doc.add_paragraph("2024年上半年全省耕地流入***万亩，流出***万亩，净增加***万亩。其中，耕地流出中，主要流向建设用地***万亩，农村道路***万亩。")
    doc.add_paragraph("全省耕地净减少的地市有*个，其中净减少最多的为温州市，净减少***万亩。")
    doc.add_paragraph("表1 各市耕地净变化情况表")
    table = doc.add_table(rows=3, cols=4)
    for index, label in enumerate(("县级行政区", "流出面积", "流入面积", "净变化")):
        table.cell(0, index).text = label
    table.cell(1, 0).text = "合计"
    table.cell(2, 0).text = "杭州市"
    doc.add_paragraph("二、疑似新增违法占耕依然发生")
    doc.add_paragraph("全省疑似新增违法建设用地***万亩。")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.mark.asyncio
async def test_client_trial_queries_only_fixed_sql() -> None:
    skills_dir = Path(__file__).resolve().parents[2] / "skills"
    result = await query_client_trial(_FakePg(), skills_dir)
    assert result["total"]["inflow"] == 40000
    assert result["cities"][0]["code"] == "3301"


def test_client_trial_fills_known_values_and_marks_unknown(tmp_path: Path) -> None:
    source = tmp_path / "模板.docx"
    source.write_bytes(_template())
    parsed = parse_docx(source.read_bytes(), source.name)
    summary = fill_client_trial_report(
        source, parsed,
        {"total": {
            "inflow": 40000, "outflow": 60000, "to_construction": 10000,
            "from_orchard": 12000, "from_forest": 8000,
            "from_construction": 2000, "from_pond": 1000, "construction": 50000,
        }, "cities": [{"code": "3301", "inflow": 10000, "outflow": 30000},
                      {"code": "3302", "inflow": 30000, "outflow": 30000}],
         "counties": [{"code": "330101", "name": "甲区", "inflow": 10000, "outflow": 30000}]},
        tmp_path / "结果.docx",
        query_started_at=datetime.fromisoformat("2026-09-27T10:00:00+08:00"),
    )
    doc = Document(tmp_path / "结果.docx")
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "2024年" not in text
    assert "60.00亩" in text
    assert "待核实" in text
    assert doc.tables[0].cell(0, 0).text == "设区市"
    assert doc.tables[0].cell(1, 0).text == "合计"
    assert summary["filled_tables"] == 1


def test_client_trial_endpoint_rejects_other_templates(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))
    app = create_app()
    with TestClient(app) as client:
        parsed = client.post("/api/report-templates/parse?filename=其他.docx", content=_template())
        assert parsed.status_code == 200
        payload = parsed.json()
        assert payload["client_trial_eligible"] is False
        response = client.post(
            f'/api/report-templates/{payload["workflow_id"]}/auto-generate-client', json={}
        )
        assert response.status_code == 409
