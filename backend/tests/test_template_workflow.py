"""标准问题库示例快报工作流回归。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from geoagent.report.question_library import load_question_library, match_parsed_questions
from geoagent.report.template import decompose_questions, parse_docx
from geoagent.report.workflow import (
    build_question_library_template,
    fill_question_library_report,
    is_question_library_template,
)
from geoagent.server.app import create_app


def _executions() -> dict:
    return {
        "total_summary": {"columns": ["n", "area"], "rows": [{"n": 100, "area": 100000.0}]},
        "orig_farmland": {
            "columns": ["n", "area_m2"], "rows": [{"n": 40, "area_m2": 40000.0}],
        },
        "current_farmland": {
            "columns": ["n", "area_m2"], "rows": [{"n": 30, "area_m2": 30000.0}],
        },
        "by_type_summary": {
            "columns": ["TBLX", "n", "area_m2"],
            "rows": [
                {"TBLX": "动土", "n": 50, "area_m2": 50000.0},
                {"TBLX": "耕地", "n": 30, "area_m2": 30000.0},
            ],
        },
        "by_county": {
            "columns": ["region", "n", "area"],
            "rows": [
                {"region": "甲县", "n": 60, "area": 60000.0},
                {"region": "乙县", "n": 40, "area": 40000.0},
            ],
        },
    }


def test_sample_template_matches_five_standard_questions(tmp_path) -> None:
    template = build_question_library_template(tmp_path / "模板.docx")
    parsed = parse_docx(template.read_bytes(), template.name)
    assert is_question_library_template(parsed)
    library = load_question_library(
        Path(__file__).resolve().parents[2]
        / "skills" / "qa-library" / "assets" / "question_library.json"
    )
    questions = match_parsed_questions(decompose_questions(parsed), library)
    assert [item["question_library_match"]["candidates"][0]["id"] for item in questions] == [
        "total_summary",
        "orig_farmland",
        "current_farmland",
        "by_type_summary",
        "by_county",
    ]
    assert [item["chart_type_hint"] for item in parsed["chart_anchors"]] == ["pie", "bar"]


def test_modified_sample_is_not_automatically_generated(tmp_path) -> None:
    template = build_question_library_template(tmp_path / "模板.docx")
    document = Document(template)
    paragraph = next(item for item in document.paragraphs if "本快报按照统一统计口径" in item.text)
    paragraph.text = paragraph.text.replace("统一统计口径", "另一统计口径")
    document.save(template)
    parsed = parse_docx(template.read_bytes(), template.name)
    assert not is_question_library_template(parsed)


def test_refill_creates_pie_and_bar_charts_without_placeholders(tmp_path) -> None:
    template = build_question_library_template(tmp_path / "模板.docx")
    parsed = parse_docx(template.read_bytes(), template.name)
    output = fill_question_library_report(
        template,
        parsed,
        _executions(),
        tmp_path / "结果.docx",
        query_started_at=datetime.fromisoformat("2026-09-21T15:00:00+08:00"),
    )
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "{{" not in text
    assert "2026-09-21 15:00" in text
    assert document.tables[0].cell(1, 0).text == "动土"
    assert document.tables[1].cell(1, 0).text == "甲县"
    with ZipFile(output) as package:
        charts = sorted(name for name in package.namelist() if name.startswith("word/charts/chart"))
        xml = "\n".join(package.read(name).decode("utf-8") for name in charts)
    assert len(charts) == 2
    assert "pieChart" in xml
    assert "barChart" in xml


class _FakePg:
    async def run_sql(self, sql: str, **_: object) -> dict:
        if 'GROUP BY t."TBLX"' in sql:
            return _executions()["by_type_summary"]
        if 'GROUP BY "XMC"' in sql:
            return _executions()["by_county"]
        if '"DLBM" LIKE' in sql:
            return _executions()["orig_farmland"]
        if '"TBLX" IN' in sql:
            return _executions()["current_farmland"]
        return _executions()["total_summary"]

    async def close(self) -> None:
        return None


def test_http_sample_auto_generates_without_question_clicks(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))

    def fake_pdf(source: Path, output: Path) -> Path:
        output.write_bytes(b"%PDF-1.4\n%%EOF")
        return output

    monkeypatch.setattr("geoagent.server.template_routes.convert_docx_to_pdf", fake_pdf)
    app = create_app()
    app.state.pg = _FakePg()
    with TestClient(app) as client:
        conversation = client.post("/api/conversations", json={"title": "自动快报"}).json()
        sample = client.get("/api/report-templates/sample").content
        parsed = client.post("/api/report-templates/parse?filename=示例.docx", content=sample)
        assert parsed.status_code == 200
        payload = parsed.json()
        assert payload["auto_eligible"] is True
        generated = client.post(
            f'/api/report-templates/{payload["workflow_id"]}/auto-generate',
            json={"conversation_id": conversation["id"]},
        )
        assert generated.status_code == 200
        files = generated.json()
        assert files["review_only"] is False
        assert client.get(files["word"]["url"]).status_code == 200
        assert client.get(files["pdf"]["url"]).status_code == 200
        messages = client.get(f'/api/conversations/{conversation["id"]}/messages').json()["messages"]
        assert [item["role"] for item in messages] == ["user", "assistant", "tool", "assistant"]

        other = Document()
        other.add_paragraph("全省变化图斑{{total_n}}个。")
        other_path = tmp_path / "非标准.docx"
        other.save(other_path)
        parsed_other = client.post(
            "/api/report-templates/parse?filename=非标准.docx", content=other_path.read_bytes()
        ).json()
        assert parsed_other["auto_eligible"] is False
        refused = client.post(
            f'/api/report-templates/{parsed_other["workflow_id"]}/auto-generate', json={}
        )
        assert refused.status_code == 409


def test_supported_endpoint_generates_library_template(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))

    def fake_pdf(source: Path, output: Path) -> Path:
        output.write_bytes(b"%PDF-1.4\n%%EOF")
        return output

    monkeypatch.setattr("geoagent.report.supported_templates.convert_docx_to_pdf", fake_pdf)
    app = create_app()
    app.state.pg = _FakePg()
    with TestClient(app) as client:
        sample = client.get("/api/report-templates/sample").content
        parsed = client.post("/api/report-templates/parse?filename=示例.docx", content=sample)
        assert parsed.status_code == 200
        payload = parsed.json()
        assert payload["supported_template_kind"] == "library"
        generated = client.post(
            f'/api/report-templates/{payload["workflow_id"]}/auto-generate-supported', json={}
        )
        assert generated.status_code == 200
        files = generated.json()
        assert files["review_only"] is True
        assert files["summary"]["charts"] == 2
        assert client.get(files["word"]["url"]).status_code == 200
        assert client.get(files["pdf"]["url"]).status_code == 200


def test_free_answer_only_binds_single_matching_value(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))

    def fake_pdf(source: Path, output: Path) -> Path:
        output.write_bytes(b"%PDF-1.4\n%%EOF")
        return output

    monkeypatch.setattr("geoagent.server.template_routes.convert_docx_to_pdf", fake_pdf)
    source = tmp_path / "自定义.docx"
    document = Document()
    document.add_paragraph("本期变化图斑{{total_n}}个。")
    document.save(source)
    app = create_app()
    with TestClient(app) as client:
        conv = client.post("/api/conversations", json={"title": "自由问数"}).json()
        parsed = client.post(
            "/api/report-templates/parse?filename=自定义.docx", content=source.read_bytes()
        ).json()
        question = parsed["questions"][0]
        url = (
            f'/api/report-templates/{parsed["workflow_id"]}/questions/'
            f'{question["id"]}/capture-free'
        )
        app.state.store.add_message(conv["id"], {"role": "user", "content": question["question"]})
        app.state.store.add_message(conv["id"], {"role": "assistant", "content": "12个，面积34亩"})
        ambiguous = client.post(url, json={"conversation_id": conv["id"]})
        assert ambiguous.status_code == 200
        assert ambiguous.json()["status"] == "needs_review"
        app.state.store.add_message(conv["id"], {"role": "user", "content": question["question"]})
        app.state.store.add_message(conv["id"], {"role": "assistant", "content": "12个"})
        bound = client.post(url, json={"conversation_id": conv["id"]})
        assert bound.status_code == 200
        assert bound.json()["status"] == "bound_reference"
        generated = client.post(
            f'/api/report-templates/{parsed["workflow_id"]}/generate',
            json={
                "conversation_id": conv["id"], "binding_mode": True,
                "confirmed_slot_ids": [bound.json()["slot_id"]],
            },
        )
        assert generated.status_code == 200
        assert generated.json()["reference_only"] is True
        output = Document(tmp_path / "reports" / generated.json()["word"]["filename"])
        text = "\n".join(paragraph.text for paragraph in output.paragraphs)
        assert "12个" in text
        assert "仅供参考" in text


def test_http_workflow_executes_each_question_and_generates_files(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))

    def fake_pdf(source: Path, output: Path) -> Path:
        output.write_bytes(b"%PDF-1.4\n%%EOF")
        return output

    monkeypatch.setattr("geoagent.server.template_routes.convert_docx_to_pdf", fake_pdf)
    app = create_app()
    app.state.pg = _FakePg()
    with TestClient(app) as client:
        conversation = client.post("/api/conversations", json={"title": "模板问数"}).json()
        template = client.get("/api/report-templates/sample")
        assert template.status_code == 200
        parsed = client.post(
            "/api/report-templates/parse?filename=示例模板.docx",
            content=template.content,
        )
        assert parsed.status_code == 200
        payload = parsed.json()
        assert len(payload["questions"]) == 5
        for question in payload["questions"]:
            library_id = question["question_library_match"]["candidates"][0]["id"]
            response = client.post(
                f'/api/report-templates/{payload["workflow_id"]}/questions/{question["id"]}/execute',
                json={
                    "library_id": library_id,
                    "question": question["question"],
                    "conversation_id": conversation["id"],
                },
            )
            assert response.status_code == 200
        messages = client.get(
            f'/api/conversations/{conversation["id"]}/messages'
        ).json()["messages"]
        assert [item["role"] for item in messages] == [
            role for _ in range(5) for role in ("user", "assistant", "tool", "assistant")
        ]
        assert all(item.get("route") == "sql" for item in messages if item["role"] == "assistant")
        generated = client.post(
            f'/api/report-templates/{payload["workflow_id"]}/generate',
            json={
                "conversation_id": conversation["id"],
                "request_text": "请生成Word和PDF快报。",
            },
        )
        assert generated.status_code == 200
        files = generated.json()
        assert files["pdf"] is not None
        assert client.get(files["word"]["url"]).status_code == 200
        assert client.get(files["pdf"]["url"]).status_code == 200
        final_messages = client.get(
            f'/api/conversations/{conversation["id"]}/messages'
        ).json()["messages"]
        assert [item["role"] for item in final_messages[-4:]] == [
            "user", "assistant", "tool", "assistant"
        ]
        assert final_messages[-2]["name"] == "generate_monitoring_report"
        assert {item["kind"] for item in final_messages[-2]["artifacts"]} == {"file"}
        assert len(final_messages[-2]["artifacts"]) == 2
