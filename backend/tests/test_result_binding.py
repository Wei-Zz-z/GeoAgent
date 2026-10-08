"""普通占位符、空表格和图题的查询结果绑定回归。"""

from datetime import datetime
import json
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

from geoagent.report.template import parse_docx, decompose_questions
from geoagent.report.result_binding import bind_result, result_column, bind_charts
from geoagent.report.generic_report import fill_generic_report
from geoagent.server.app import create_app
from geoagent.report.reference_answer import query_scalar
from geoagent.report.question_library import load_question_library, match_question
from pathlib import Path


def make_template(path):
    document = Document()
    document.add_paragraph("查询时间：[[SYSTEM_TIME]]")
    document.add_paragraph("本期全库共有***个变化图斑，总面积***亩。")
    document.add_paragraph("表1 各县变化图斑数量和总面积")
    document.add_paragraph("单位：亩")
    table = document.add_table(rows=3, cols=3)
    for cell, label in zip(table.rows[0].cells, ["县（市、区）", "图斑数量", "面积"]):
        cell.text = label
    document.add_paragraph("图1 各县变化图斑面积对比")
    document.add_paragraph("图2 各县变化图斑面积构成")
    document.add_paragraph("指定面积区间的平均图斑面积***平方米。")
    document.save(path)
    return parse_docx(path.read_bytes(), path.name)


def test_bind_plain_placeholders_and_table_by_columns(tmp_path):
    source = tmp_path / "template.docx"
    parsed = make_template(source)
    questions = decompose_questions(parsed)
    scalar = bind_result(questions[0], parsed, {"rows": [{"n": 12, "area": 1000}]})
    assert list(scalar["scalars"].values()) == ["12", "1.50"]
    table = bind_result(questions[1], parsed, {"rows": [{"region": "甲县", "n": 2, "area": 1000}]})
    assert list(table["tables"].values()) == [[["甲县", "2", "1.50"]]]
    assert result_column("面积", ["area", "area_m2"]) is None
    assert result_column("耕地流入面积", ["n", "area"]) is None
    assert bind_result(questions[0], parsed, {"rows": [{"n": 1}, {"n": 2}]})["scalars"] == {}
    assert bind_result(questions[0], parsed, {"rows": [{"n": 1, "area": None}]})["scalars"] == {"b1:s0": "1"}
    assert bind_result(questions[0], parsed, {"rows": [{"n": 1, "area": 1000}], "truncated": True})["scalars"] == {}


def test_generic_output_shares_table_chart_data_and_marks_free_result(tmp_path):
    source = tmp_path / "template.docx"
    parsed = make_template(source)
    questions = decompose_questions(parsed)
    scalar = bind_result(questions[0], parsed, {"rows": [{"n": 12, "area": 1000}]})
    bound = {key: {"value": value, "status": "标准问题自动绑定"} for key, value in scalar["scalars"].items()}
    bound[questions[-1]["slot_ids"][0]] = {"value": "123.45", "status": "仅供参考"}
    table = bind_result(questions[1], parsed, {"rows": [{"region": "甲县", "n": 2, "area": 1000}]})
    tables = {key: {"slot_id": key, "rows": rows, "source": table["sources"][key],
                    "source_question_id": questions[1]["id"]} for key, rows in table["tables"].items()}
    output = tmp_path / "out.docx"
    summary = fill_generic_report(source, parsed, output, bound, tables, datetime.now().astimezone())
    assert summary["charts"] == 2
    assert summary["pending_items"] == []
    document = Document(output)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "12个变化图斑，总面积1.50亩" in text
    assert "123.45平方米" in text
    assert "未命中标准问题库，仅供参考" in text
    assert 'FFF2CC' in document._element.xml
    assert document.tables[0].alignment == WD_TABLE_ALIGNMENT.CENTER
    assert document.tables[0].cell(0, 0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.tables[0].cell(0, 0).paragraphs[0].paragraph_format.keep_with_next
    assert document.tables[0].cell(1, 0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert document.tables[0].cell(1, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert document.tables[0].cell(1, 1).vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in document.paragraphs:
        if paragraph._p.xpath(".//w:drawing") or paragraph.text.startswith("图"):
            assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    with ZipFile(output) as package:
        assert len([name for name in package.namelist() if name.startswith("word/charts/chart")]) == 2


def test_generic_endpoint_skips_pdf_and_special_template_generator(tmp_path, monkeypatch):
    monkeypatch.setenv("GEOAGENT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("GEOAGENT_REPORTS_DIR", str(tmp_path / "reports"))
    def forbidden(*args, **kwargs):
        raise AssertionError("不应调用固定模板或PDF转换")
    monkeypatch.setattr("geoagent.server.template_routes.convert_docx_to_pdf", forbidden)
    monkeypatch.setattr("geoagent.server.template_routes.fill_question_library_report", forbidden)
    source = tmp_path / "template.docx"
    make_template(source)
    app = create_app()
    with TestClient(app) as client:
        response = client.post("/api/report-templates/parse?filename=template.docx", content=source.read_bytes())
        assert response.status_code == 200
        data = response.json()
        assert data["supported_template_kind"] is None
        response = client.post(f"/api/report-templates/{data['workflow_id']}/generate",
                               json={"generic_mode": True, "include_pdf": False})
        assert response.status_code == 200, response.text
        assert response.json()["pdf"] is None
        assert response.json()["summary"]["pending_items"]


def test_free_query_scalar_uses_actual_sql_result_and_rejects_ambiguity():
    messages = [
        {"role": "assistant", "tool_calls": [{"id": "sql1", "function": {
            "name": "run_sql", "arguments": '{"sql":"SELECT AVG(\\"MJ\\") AS mean_area_m2 FROM data.\\"2026_1_change_landuse\\""}'}}]},
        {"role": "tool", "tool_call_id": "sql1", "artifacts": [{"name": "query_result", "kind": "table",
         "data": {"rows": [{"mean_area_m2": 1000}]}}]},
    ]
    assert query_scalar(messages, "亩")["value"] == "1.50"
    assert query_scalar(messages, "%") is None
    assert query_scalar(messages + [messages[-1]], "亩") is None
    assert query_scalar([], "亩") is None
    messages[0]["tool_calls"][0]["function"]["arguments"] = json.dumps({
        "sql": 'SELECT AVG("MJ") AS avg_mj FROM data."2026_1_change_landuse"'})
    messages[1]["artifacts"][0]["data"]["rows"] = [{"avg_mj": 124.026}]
    assert query_scalar(messages, "平方米")["value"] == "124.03"
    messages[0]["tool_calls"][0]["function"]["arguments"] = json.dumps({
        "sql": 'SELECT AVG("MJ") * 0.0015 AS avg_mj FROM data."2026_1_change_landuse"'})
    assert query_scalar(messages, "平方米") is None


def test_missing_average_or_median_cannot_auto_use_total_area():
    library = load_question_library(Path(__file__).resolve().parents[2] / "skills/qa-library/assets/question_library.json")
    for query in ("变化图斑平均面积是多少", "变化图斑面积中位数是多少"):
        result = match_question(query, library)
        assert result["status"] == "unmatched"
        assert all(not item["compatible"] for item in result["candidates"])


def test_pie_groups_long_tail_without_losing_total(tmp_path):
    parsed = make_template(tmp_path / "template.docx")
    slot = next(item for item in parsed["slots"] if item["kind"] == "table_region")
    tables = {slot["id"]: {
        "slot_id": slot["id"], "source_question_id": "Q2",
        "source": {"columns": ["region", "n", "area"]},
        "rows": [[f"地区{i}", "1", str(i)] for i in range(1, 9)],
    }}
    charts = bind_charts(parsed, tables)
    pie = next(chart for chart in charts.values() if chart["chart_type"] == "pie")
    assert len(pie["labels"]) == 6
    assert pie["labels"][-1] == "其余类型（合计）"
    assert sum(pie["values"]) == 36
    assert len(pie["grouped_categories"]) == 3
